"""
datasphere_profiles.py — perfiles de conexión Datasphere por CLIENTE para el dueño de la app (Assessment ad-hoc + Publicar modelo).

Fuente ÚNICA: el libro Excel privado de SharePoint que apunta ``config.SHAREPOINT_DSP_OAUTH_URL`` (link de compartir). Se lee con el
MISMO camino que el resto de la integración (`generate_catalog.fetch_shared_file`: Graph ``/shares`` con el token DELEGADO del usuario, el
que el navegador ya envía como ``graph_token``). No hay cliente SharePoint nuevo, ni URL en el código, ni otra variable de entorno.

ESQUEMA (observado en el libro de perfiles; hoja ``OAuth``): una fila de encabezado con ``Cliente | URL Host | URL Token | Client ID |
Secret`` y una fila por cliente debajo (las filas superiores, ``Name`` / ``User ID``, son metadatos y se ignoran). El encabezado se busca por
NOMBRE (no por posición) y se exigen las cinco columnas: si falta alguna el libro se considera inválido (no se inventan columnas). El libro NO
trae "Space": ese campo sigue siendo manual.

SEGURIDAD (contrato central):
  * el navegador SOLO recibe metadatos seguros (``public()``: profile_id, cliente, host, token URL, client ID); el ``secret`` vive únicamente
    en el servidor y se vuelve a resolver desde SharePoint en CADA ejecución a partir del ``profile_id`` opaco (nunca se cachea ni se guarda);
  * el libro se baja a un directorio temporal que se elimina enseguida (nunca a disco persistido/espejado) y su contenido no se registra;
  * todo es OPCIONAL: sin URL configurada, sin token, SharePoint caído o libro inválido → estado estable y el formulario manual sigue igual.
"""
from __future__ import annotations

import hashlib
import logging
import re
import sys
import tempfile
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path

import config

log = logging.getLogger("danwor.datasphere_profiles")

# Estados estables (el navegador decide la copia; nunca recibe detalle del error de SharePoint).
OK, NOT_CONFIGURED, NO_TOKEN, UNAVAILABLE, INVALID_WORKBOOK, EMPTY = (
    "ok", "not_configured", "no_token", "unavailable", "invalid_workbook", "empty")

WARNINGS = {
    NOT_CONFIGURED: "La fuente de perfiles de clientes no está configurada. Completá la conexión a mano.",
    NO_TOKEN: "No se pudo autorizar el acceso a SharePoint. Reingresá o completá la conexión a mano.",
    UNAVAILABLE: "No se pudo leer la lista de clientes desde SharePoint. Completá la conexión a mano.",
    INVALID_WORKBOOK: "El libro de clientes no tiene el formato esperado. Completá la conexión a mano.",
    EMPTY: "No hay clientes disponibles en la lista. Completá la conexión a mano.",
}
RUN_FAILURE_TITLE = "No se pudo usar el perfil del cliente"
RUN_FAILURE_MESSAGE = ("No se pudo resolver el perfil del cliente desde SharePoint. Volvé a elegirlo o completá la conexión a mano.")

# Encabezados reconocidos (normalizados). Solo variantes de los nombres observados: no se inventan columnas.
_HEADERS = {
    "client": {"cliente", "client", "nombre de cliente"},
    "host": {"url host", "host", "url del tenant", "url tenant"},
    "token_url": {"url token", "token url", "token url oauth"},
    "client_id": {"client id", "clientid"},
    "secret": {"secret", "client secret", "clientsecret"},
}


@dataclass(frozen=True)
class Profile:
    client: str
    host: str
    token_url: str
    client_id: str
    secret: str = field(repr=False)                     # JAMÁS sale del servidor

    @property
    def profile_id(self) -> str:
        return profile_id_for(self.client)

    def public(self) -> dict:
        """Metadatos SEGUROS para el navegador. Sin el secreto."""
        return {"profile_id": self.profile_id, "client": self.client, "host": self.host,
                "token_url": self.token_url, "client_id": self.client_id}


def profile_id_for(client: str) -> str:
    """Identificador opaco y determinístico del perfil (derivado del nombre del cliente): permite volver a resolverlo del libro sin
    guardar nada en el servidor. No contiene el secreto ni lo revela."""
    return hashlib.sha256(("dsp-owner-profile\0" + _norm(client)).encode("utf-8")).hexdigest()[:20]


def source_url() -> str:
    return (getattr(config, "SHAREPOINT_DSP_OAUTH_URL", "") or "").strip()


def source_configured() -> bool:
    return bool(source_url())


def _norm(value) -> str:
    text = unicodedata.normalize("NFKD", str(value if value is not None else ""))
    text = "".join(c for c in text if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", text).strip().casefold()


# ───────────────────────────── lectura del libro ─────────────────────────────
def _download_workbook(token: str, dest_dir: Path) -> Path:
    """Baja el libro con la integración EXISTENTE (`fetch_shared_file`) a ``dest_dir`` (temporal) y devuelve su ruta."""
    tools_catalog = str(Path(__file__).resolve().parents[2] / "tools" / "catalog")
    if tools_catalog not in sys.path:                     # igual que `admin_tools.run_catalog_generation`: el generador importa módulos hermanos
        sys.path.insert(0, tools_catalog)
    import generate_catalog as gen
    gen.fetch_shared_file(source_url(), token, dest_dir, fallback_name="perfiles-datasphere.xlsx")
    files = [p for p in Path(dest_dir).iterdir() if p.is_file()]
    if not files:
        raise FileNotFoundError("sin archivo descargado")
    return files[0]


def parse_workbook(path: Path) -> list[Profile]:
    """Perfiles del libro, en orden. Busca el encabezado por nombre en cada hoja; ``ValueError`` si ninguna hoja lo tiene completo.
    Las filas incompletas se omiten y un cliente repetido conserva solo la primera aparición."""
    from openpyxl import load_workbook
    wb = load_workbook(str(path), read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            rows = list(ws.iter_rows(values_only=True))
            for r_index, row in enumerate(rows[:50]):
                cols = _header_columns(row)
                if cols is None:
                    continue
                return _profiles_from_rows(rows[r_index + 1:], cols)
    finally:
        wb.close()
    raise ValueError("libro sin encabezado de perfiles")


def _header_columns(row) -> dict[str, int] | None:
    found: dict[str, int] = {}
    for index, cell in enumerate(row or ()):
        name = _norm(cell)
        for key, aliases in _HEADERS.items():
            if name in aliases and key not in found:
                found[key] = index
    return found if set(found) == set(_HEADERS) else None


def _profiles_from_rows(rows, cols: dict[str, int]) -> list[Profile]:
    out: list[Profile] = []
    seen: set[str] = set()

    def cell(row, key: str) -> str:
        i = cols[key]
        return str(row[i]).strip() if i < len(row) and row[i] is not None else ""

    for row in rows:
        values = {k: cell(row, k) for k in cols}
        if not all(values.values()):                      # fila vacía o incompleta
            continue
        ident = _norm(values["client"])
        if ident in seen:
            continue
        seen.add(ident)
        out.append(Profile(values["client"], values["host"], values["token_url"], values["client_id"], values["secret"]))
    return out


def load_profiles(graph_token: str) -> tuple[str, list[Profile]]:
    """(estado, perfiles). NUNCA lanza. Lee SharePoint cada vez: no hay caché ni copia local del libro (contiene secretos)."""
    if not source_configured():
        return NOT_CONFIGURED, []
    if not (graph_token or "").strip():
        return NO_TOKEN, []
    try:
        with tempfile.TemporaryDirectory(prefix="dsp-profiles-") as tmp:
            path = _download_workbook(graph_token.strip(), Path(tmp))
            profiles = parse_workbook(path)
    except ValueError:
        log.warning("datasphere_profiles: el libro de perfiles no tiene el formato esperado")
        return INVALID_WORKBOOK, []
    except Exception as exc:                              # noqa: BLE001 — jamás se registra el contenido; solo el tipo de error
        log.warning("datasphere_profiles: no se pudo leer la fuente de perfiles (%s)", type(exc).__name__)
        return UNAVAILABLE, []
    return (OK, profiles) if profiles else (EMPTY, [])


# ───────────────────────────── API para las rutas ─────────────────────────────
def list_public(graph_token: str) -> dict:
    """Respuesta SEGURA para el selector del navegador: ``{"ok", "profiles": [...]}`` o ``{"ok": False, "reason", "message"}``."""
    status, profiles = load_profiles(graph_token)
    if status != OK:
        return {"ok": False, "reason": status, "message": WARNINGS[status]}
    return {"ok": True, "profiles": [p.public() for p in profiles]}


def resolve(profile_id: str, graph_token: str) -> Profile | None:
    """Perfil completo (CON secreto) para uso EXCLUSIVO del servidor. ``None`` si no se pudo resolver."""
    wanted = (profile_id or "").strip()
    if not wanted:
        return None
    status, profiles = load_profiles(graph_token)
    if status != OK:
        return None
    return next((p for p in profiles if p.profile_id == wanted), None)


def selector_url(stem: str, can_execute: bool) -> str:
    """URL del endpoint de perfiles si el selector debe renderizarse (dueño + fuente configurada); si no, ''. Decide el SERVIDOR."""
    return f"/modules/{stem}/datasphere-profiles" if (can_execute and source_configured()) else ""


def failed_run_events() -> str:
    """Cuerpo NDJSON de una ejecución que no pudo resolver su perfil (mismo contrato que el stream del run; sin ejecutar nada)."""
    import json
    start = json.dumps({"t": "start"}) + "\n"
    done = json.dumps({"t": "done", "state": "FAILED", "error_code": "PROFILE_UNAVAILABLE", "message": RUN_FAILURE_MESSAGE,
                       "title": RUN_FAILURE_TITLE, "errors": {}, "operation_id": None, "duration_s": 0, "result_url": None},
                      ensure_ascii=False) + "\n"
    return start + done
