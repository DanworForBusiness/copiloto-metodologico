"""
test_profile_import_service.py — Capital Humano Phase 3.8 (Safe Bulk
Profile Import): contrato de `webapp.services.profile_import_service` —
parseo XLSX, preview (cero writes), integridad de token firmado, y
confirm — sin HTTP (ver test_human_capital_profile_import_routes.py para
eso), sin Spaces real."""
from __future__ import annotations

import io
import zipfile

import openpyxl
import pytest
from jose import jwt

import config
from webapp.services import human_capital_service, profile_repository, profile_import_service as pis
from webapp.services import human_capital_catalog_repository as catalog_repo


@pytest.fixture(autouse=True)
def _isolated_base_dir(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "BASE_DIR", tmp_path)
    monkeypatch.setattr(profile_repository.remote_store, "push_path", lambda path: True)
    monkeypatch.setattr(profile_repository.remote_store, "enabled", lambda: False)
    # Fuente migró a catálogo GOBERNADO (CCHH Catalog Foundation) -- import
    # sigue rechazando lo que no está en catálogo ("no cambiar semántica
    # todavía", ver validate_source_for_import), así que este archivo
    # necesita sembrar el ÚNICO valor no-vacío que usa para una fila
    # eligible ("Portal de empleo" es el caso de rechazo deliberado).
    snap = catalog_repo.build_snapshot(
        {"fuente": [{"valor": "Flyer", "orden": 1}]},
        refreshed_by="test", source={"drive_id": "d", "item_id": "i"},
    )
    p = catalog_repo.snapshot_path()
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_bytes(catalog_repo._serialize(snap))


def _workbook_bytes(headers: list[str], rows: list[list]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


_ALL_HEADERS = ["nombre", "apellido", "email", "telefono", "linkedin", "ubicacion",
                "perfil_skills", "seniority", "fuente", "tags", "notas", "potencial", "marca_potencial"]


def _row(nombre="Juan", apellido="Perez", email="", fuente="", potencial="") -> list:
    return [nombre, apellido, email, "", "", "", "", "", fuente, "", "", potencial, ""]


# ============================================================================
# Upload / parsing (§28 "Upload")
# ============================================================================

def test_valid_xlsx_parses_successfully():
    data = _workbook_bytes(_ALL_HEADERS, [_row()])
    rows = pis.parse_workbook(data)
    assert len(rows) == 1
    assert rows[0]["nombre"] == "Juan"


def test_empty_workbook_raises_import_file_error():
    wb = openpyxl.Workbook()
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(pis.ImportFileError):
        pis.parse_workbook(buf.getvalue())


def test_invalid_workbook_bytes_raise_import_file_error_not_a_raw_traceback():
    with pytest.raises(pis.ImportFileError):
        pis.parse_workbook(b"esto no es un xlsx valido")


def test_missing_required_headers_raise_import_file_error():
    data = _workbook_bytes(["email", "telefono"], [["a@x.com", "123"]])
    with pytest.raises(pis.ImportFileError):
        pis.parse_workbook(data)


def test_row_limit_is_enforced(monkeypatch):
    monkeypatch.setattr(pis, "MAX_IMPORT_ROWS", 3)
    data = _workbook_bytes(_ALL_HEADERS, [_row() for _ in range(5)])
    with pytest.raises(pis.ImportFileError):
        pis.parse_workbook(data)


def test_blank_rows_are_skipped_and_do_not_count_toward_the_limit(monkeypatch):
    monkeypatch.setattr(pis, "MAX_IMPORT_ROWS", 2)
    data = _workbook_bytes(_ALL_HEADERS, [_row(), [None] * len(_ALL_HEADERS), _row(nombre="Ana")])
    rows = pis.parse_workbook(data)
    assert len(rows) == 2


def test_optional_columns_may_be_entirely_absent():
    data = _workbook_bytes(["nombre", "apellido"], [["Juan", "Perez"]])
    rows = pis.parse_workbook(data)
    assert len(rows) == 1
    preview = pis.build_preview(rows)
    assert preview[0].status == "eligible"


# ============================================================================
# Row normalization (§8 del ticket 3.8) — reusa validate_and_shape_fields,
# nunca un validador bulk paralelo
# ============================================================================

def test_flyer_source_is_accepted_via_bulk_exactly_like_manual():
    data = _workbook_bytes(_ALL_HEADERS, [_row(fuente="Flyer")])
    preview = pis.build_preview(pis.parse_workbook(data))
    assert preview[0].status == "eligible"
    assert preview[0].fields["source"] == "Flyer"


def test_fractional_potential_rating_is_rejected_not_silently_truncated():
    data = _workbook_bytes(_ALL_HEADERS, [_row(potencial="1.5")])
    preview = pis.build_preview(pis.parse_workbook(data))
    assert preview[0].status == "invalid"


def test_integral_float_potential_rating_cell_is_accepted():
    # Una celda XLSX que guarda "1" como float 1.0 (típico de Excel) debe
    # seguir siendo válida -- _cell_to_str la normaliza a "1", nunca a
    # algo que int() rechace.
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(_ALL_HEADERS)
    ws.append(["Juan", "Perez", "", "", "", "", "", "", "", "", "", 1.0, ""])
    buf = io.BytesIO()
    wb.save(buf)
    preview = pis.build_preview(pis.parse_workbook(buf.getvalue()))
    assert preview[0].status == "eligible"
    assert preview[0].fields["potential_rating"] == 1


def test_invalid_source_is_rejected():
    data = _workbook_bytes(_ALL_HEADERS, [_row(fuente="Portal de empleo")])
    preview = pis.build_preview(pis.parse_workbook(data))
    assert preview[0].status == "invalid"


def test_missing_required_field_within_row_is_rejected():
    data = _workbook_bytes(_ALL_HEADERS, [_row(nombre="")])
    preview = pis.build_preview(pis.parse_workbook(data))
    assert preview[0].status == "invalid"


# ============================================================================
# Preview — zero writes, list_profiles() exactamente una vez (§9/§11)
# ============================================================================

def test_preview_performs_zero_persistence_calls(monkeypatch):
    calls = {"create": 0, "create_explicit": 0, "push": 0}
    monkeypatch.setattr(profile_repository, "create", lambda *a, **k: calls.__setitem__("create", calls["create"] + 1))
    monkeypatch.setattr(profile_repository, "create_with_explicit_id", lambda *a, **k: calls.__setitem__("create_explicit", calls["create_explicit"] + 1))
    monkeypatch.setattr(profile_repository.remote_store, "push_path", lambda p: calls.__setitem__("push", calls["push"] + 1) or True)
    data = _workbook_bytes(_ALL_HEADERS, [_row(), _row(nombre="Ana", apellido="Gomez")])
    pis.build_preview(pis.parse_workbook(data))
    assert calls == {"create": 0, "create_explicit": 0, "push": 0}


def test_preview_calls_list_profiles_exactly_once(monkeypatch):
    calls = []
    original = human_capital_service.list_profiles
    monkeypatch.setattr(human_capital_service, "list_profiles", lambda: (calls.append(1), original())[1])
    data = _workbook_bytes(_ALL_HEADERS, [_row(), _row(nombre="Ana", apellido="Gomez"), _row(nombre="Luis", apellido="Diaz")])
    pis.build_preview(pis.parse_workbook(data))
    assert len(calls) == 1


def test_eligible_row_has_no_existing_profiles():
    data = _workbook_bytes(_ALL_HEADERS, [_row()])
    preview = pis.build_preview(pis.parse_workbook(data))
    assert preview[0].status == "eligible"
    assert preview[0].token is not None
    assert preview[0].profile_id is not None


def test_row_matching_an_existing_profile_is_flagged():
    human_capital_service.create_profile({"first_name": "Juan", "last_name": "Perez", "email": "juan@x.com"})
    data = _workbook_bytes(_ALL_HEADERS, [_row(email="juan@x.com")])
    preview = pis.build_preview(pis.parse_workbook(data))
    assert preview[0].status == "existing_duplicate"
    assert preview[0].token is not None  # sigue elegible, solo con advisory


def test_within_file_duplicate_flags_the_second_occurrence_against_the_first():
    data = _workbook_bytes(_ALL_HEADERS, [
        _row(email="juan@x.com"),
        _row(email="JUAN@x.com"),
    ])
    preview = pis.build_preview(pis.parse_workbook(data))
    assert preview[0].status == "eligible"
    assert preview[1].status == "within_file_duplicate"
    assert preview[1].duplicate_row_number == 1


def test_a_third_matching_row_still_anchors_to_the_first_occurrence():
    data = _workbook_bytes(_ALL_HEADERS, [
        _row(email="juan@x.com"),
        _row(email="juan@x.com"),
        _row(email="juan@x.com"),
    ])
    preview = pis.build_preview(pis.parse_workbook(data))
    assert preview[1].duplicate_row_number == 1
    assert preview[2].duplicate_row_number == 1


def test_both_within_file_duplicate_rows_remain_visible_in_preview():
    data = _workbook_bytes(_ALL_HEADERS, [_row(email="juan@x.com"), _row(email="juan@x.com")])
    preview = pis.build_preview(pis.parse_workbook(data))
    assert len(preview) == 2
    assert all(p.status in ("eligible", "within_file_duplicate") for p in preview)


def test_invalid_row_never_gets_a_profile_id_or_token():
    data = _workbook_bytes(_ALL_HEADERS, [_row(nombre="")])
    preview = pis.build_preview(pis.parse_workbook(data))
    assert preview[0].profile_id is None
    assert preview[0].token is None


# ============================================================================
# Token integrity (§13/§24) — nunca confiar en el estado del browser sin
# verificar la firma primero
# ============================================================================

def test_valid_signed_token_verifies():
    data = _workbook_bytes(_ALL_HEADERS, [_row()])
    preview = pis.build_preview(pis.parse_workbook(data))
    claims = pis.verify_row_token(preview[0].token)
    assert claims["profile_id"] == preview[0].profile_id
    assert claims["fields"]["first_name"] == "Juan"


_B64 = "ABCDEFGHIJKLMNOPQRSTUVWXYZabcdefghijklmnopqrstuvwxyz0123456789-_"


def _tokens(n):
    """n tokens legítimos distintos (el `exp` cambia y con él la firma): la prueba no depende de UNA firma concreta."""
    claims = {"purpose": pis._ROW_TOKEN_PURPOSE, "profile_id": "p", "row_number": 1, "fields": {"first_name": "Juan"}, "has_duplicate_advisory": False}
    return [jwt.encode({**claims, "exp": 1_900_000_000 + i}, config.SESSION_SECRET, algorithm="HS256") for i in range(n)]


def test_tampered_token_is_rejected():
    data = _workbook_bytes(_ALL_HEADERS, [_row()])
    preview = pis.build_preview(pis.parse_workbook(data))
    head, payload, sig = preview[0].token.split(".")
    # Mutación que cambia inequívocamente bytes FIRMADOS: el PRIMER carácter de la firma lleva 6 bits significativos (el último solo 4 de 6).
    tampered = ".".join([head, payload, ("B" if sig[0] != "B" else "C") + sig[1:]])
    with pytest.raises(pis.RowTokenError):
        pis.verify_row_token(tampered)
    # y el payload: cualquier cambio invalida la firma
    flipped = ".".join([head, payload[:-2] + ("AA" if payload[-2:] != "AA" else "BB"), sig])
    with pytest.raises(pis.RowTokenError):
        pis.verify_row_token(flipped)


def test_tampering_is_rejected_for_many_distinct_signatures():
    """Regresión de la intermitencia: la mutación «últimos dos caracteres → aa» dejaba la MISMA firma en ~0,1 % de los tokens (el último carácter base64url de una firma de
    32 bytes solo usa 4 de sus 6 bits). Con la mutación sobre bytes firmados no hay token que se acepte."""
    for tok in _tokens(400):
        head, payload, sig = tok.split(".")
        bad = ".".join([head, payload, ("B" if sig[0] != "B" else "C") + sig[1:]])
        with pytest.raises(pis.RowTokenError):
            pis.verify_row_token(bad)
        assert pis.verify_row_token(tok)["purpose"] == pis._ROW_TOKEN_PURPOSE         # y el legítimo sigue verificando


def test_non_canonical_base64_representations_of_a_valid_signature_are_rejected():
    for tok in _tokens(60):
        head, payload, sig = tok.split(".")
        last = sig[-1]
        same_bits = [c for c in _B64 if (_B64.index(c) >> 2) == (_B64.index(last) >> 2) and c != last]    # 3 caracteres distintos que decodifican a la MISMA firma
        assert len(same_bits) == 3
        for c in same_bits:
            with pytest.raises(pis.RowTokenError):
                pis.verify_row_token(".".join([head, payload, sig[:-1] + c]))
        for noisy in (sig + "=", sig + "==", sig[:10] + "*" + sig[10:], sig[:5] + "\n" + sig[5:]):          # relleno y caracteres fuera del alfabeto que el decodificador descarta
            with pytest.raises(pis.RowTokenError):
                pis.verify_row_token(".".join([head, payload, noisy]))
        assert pis.verify_row_token(tok)["profile_id"] == "p"                            # la forma canónica es la única aceptada


def test_malformed_token_shapes_are_rejected_without_leaking_library_errors():
    for bad in ("", "a.b", "a.b.", "a.b.c.d", ".."):
        with pytest.raises(pis.RowTokenError):
            pis.verify_row_token(bad)


def test_swapped_profile_id_inside_the_token_is_rejected():
    # Decodifica SIN verificar (alg=none no soportado por jose.jwt.decode
    # sin algorithms explícito -- acá directamente re-firmamos con una
    # key DISTINTA para simular un intento de forjar el claim).
    forged = jwt.encode({"profile_id": "attacker-chosen-id", "row_number": 1, "fields": {}, "has_duplicate_advisory": False, "exp": 9999999999}, "wrong-secret", algorithm="HS256")
    with pytest.raises(pis.RowTokenError):
        pis.verify_row_token(forged)


def test_changed_payload_after_signing_is_rejected():
    data = _workbook_bytes(_ALL_HEADERS, [_row()])
    preview = pis.build_preview(pis.parse_workbook(data))
    token = preview[0].token
    # Cualquier alteración de un solo caracter en el cuerpo del JWT
    # invalida la firma -- probamos alterando un caracter del payload
    # (segunda sección, separada por ".").
    header, payload, sig = token.split(".")
    altered_payload = payload[:-1] + ("A" if payload[-1] != "A" else "B")
    tampered_token = f"{header}.{altered_payload}.{sig}"
    with pytest.raises(pis.RowTokenError):
        pis.verify_row_token(tampered_token)


def test_expired_token_is_rejected(monkeypatch):
    import time
    monkeypatch.setattr(pis, "_ROW_TOKEN_MAX_AGE_S", -1)  # ya expirado al firmarlo
    data = _workbook_bytes(_ALL_HEADERS, [_row()])
    preview = pis.build_preview(pis.parse_workbook(data))
    with pytest.raises(pis.RowTokenError):
        pis.verify_row_token(preview[0].token)


# ============================================================================
# Idempotencia / confirm (§28 "Idempotency"/"Collision"/"Unknown")
# ============================================================================

def test_first_confirm_creates_the_profile():
    data = _workbook_bytes(_ALL_HEADERS, [_row()])
    preview = pis.build_preview(pis.parse_workbook(data))
    token = preview[0].token
    results = pis.confirm_rows([token], [token])
    assert results[0].status == "CREATED"
    assert len(profile_repository.list()) == 1


def test_same_signed_confirm_again_is_already_committed_with_no_rewrite():
    data = _workbook_bytes(_ALL_HEADERS, [_row()])
    preview = pis.build_preview(pis.parse_workbook(data))
    token = preview[0].token
    first = pis.confirm_rows([token], [token])[0]
    assert first.status == "CREATED"
    saved_before = profile_repository.get(preview[0].profile_id)

    second = pis.confirm_rows([token], [token])[0]
    assert second.status == "ALREADY_COMMITTED"
    saved_after = profile_repository.get(preview[0].profile_id)
    assert saved_after["created_at"] == saved_before["created_at"]
    assert saved_after["updated_at"] == saved_before["updated_at"]
    assert saved_after["revision"] == saved_before["revision"]
    assert len(profile_repository.list()) == 1


def test_unincluded_token_is_reported_as_duplicate_skipped_without_persisting():
    data = _workbook_bytes(_ALL_HEADERS, [_row()])
    preview = pis.build_preview(pis.parse_workbook(data))
    token = preview[0].token
    results = pis.confirm_rows([token], [])  # presente en all_tokens, ausente de included
    assert results[0].status == "DUPLICATE_SKIPPED"
    assert profile_repository.list() == []


def test_unknown_outcome_preserves_the_same_token_for_retry(monkeypatch):
    data = _workbook_bytes(_ALL_HEADERS, [_row()])
    preview = pis.build_preview(pis.parse_workbook(data))
    token = preview[0].token
    monkeypatch.setattr(profile_repository.remote_store, "push_path", lambda path: False)
    monkeypatch.setattr(profile_repository.remote_store, "enabled", lambda: True)
    results = pis.confirm_rows([token], [token])
    assert results[0].status == "REMOTE_OUTCOME_UNKNOWN"
    assert results[0].token == token  # MISMO token, para reintentar esta fila puntual


def test_retrying_an_unknown_row_with_the_same_token_never_mints_a_new_uuid(monkeypatch):
    data = _workbook_bytes(_ALL_HEADERS, [_row()])
    preview = pis.build_preview(pis.parse_workbook(data))
    token = preview[0].token
    profile_id_before = preview[0].profile_id

    monkeypatch.setattr(profile_repository.remote_store, "push_path", lambda path: False)
    monkeypatch.setattr(profile_repository.remote_store, "enabled", lambda: True)
    first = pis.confirm_rows([token], [token])[0]
    assert first.status == "REMOTE_OUTCOME_UNKNOWN"

    monkeypatch.setattr(profile_repository.remote_store, "push_path", lambda path: True)
    second = pis.confirm_rows([first.token], [first.token])[0]
    assert second.status == "CREATED"
    saved = profile_repository.list()
    assert len(saved) == 1
    assert saved[0]["profile_id"] == profile_id_before


def test_invalid_token_is_reported_without_raising():
    results = pis.confirm_rows(["not-a-real-token"], ["not-a-real-token"])
    assert results[0].status == "INVALID"


def test_no_batch_stop_when_one_row_is_unknown(monkeypatch):
    data = _workbook_bytes(_ALL_HEADERS, [_row(nombre="Uno"), _row(nombre="Dos", apellido="B")])
    preview = pis.build_preview(pis.parse_workbook(data))
    tokens = [p.token for p in preview]

    calls = {"n": 0}

    def _push(path):
        calls["n"] += 1
        return calls["n"] != 1  # solo la 1ra fila falla (y su reintento -- ver abajo)

    # Fuerza a que la PRIMERA fila falle en ambos intentos (push+retry) y
    # la segunda tenga éxito -- confirma que el batch sigue procesando la
    # fila independiente restante.
    monkeypatch.setattr(profile_repository.remote_store, "enabled", lambda: True)
    fail_first_row_calls = {"n": 0}

    def _push_fail_first_row_twice(path):
        fail_first_row_calls["n"] += 1
        return fail_first_row_calls["n"] > 2

    monkeypatch.setattr(profile_repository.remote_store, "push_path", _push_fail_first_row_twice)
    results = pis.confirm_rows(tokens, tokens)
    statuses = [r.status for r in results]
    assert "REMOTE_OUTCOME_UNKNOWN" in statuses
    assert "CREATED" in statuses


def test_summarize_counts_each_status():
    data = _workbook_bytes(_ALL_HEADERS, [_row()])
    preview = pis.build_preview(pis.parse_workbook(data))
    token = preview[0].token
    results = pis.confirm_rows([token], [token])
    summary = pis.summarize(results, total_rows=1, invalid_count=0)
    assert summary["CREATED"] == 1
    assert summary["total_rows"] == 1


# ============================================================================
# Phase 3.8.1 §3/§4 — separación de dominio del token firmado
# ============================================================================

def test_row_token_carries_the_purpose_claim():
    data = _workbook_bytes(_ALL_HEADERS, [_row()])
    preview = pis.build_preview(pis.parse_workbook(data))
    claims = jwt.decode(preview[0].token, config.SESSION_SECRET, algorithms=["HS256"])
    assert claims["purpose"] == "human_capital_profile_import_row"


def test_valid_purpose_and_exp_is_accepted():
    data = _workbook_bytes(_ALL_HEADERS, [_row()])
    preview = pis.build_preview(pis.parse_workbook(data))
    assert pis.verify_row_token(preview[0].token)["profile_id"] == preview[0].profile_id


def test_wrong_purpose_is_rejected_even_with_a_valid_signature():
    forged = jwt.encode(
        {"purpose": "something_else", "profile_id": "x", "row_number": 1, "fields": {},
         "has_duplicate_advisory": False, "exp": 9999999999},
        config.SESSION_SECRET, algorithm="HS256",
    )
    with pytest.raises(pis.RowTokenError):
        pis.verify_row_token(forged)


def test_missing_purpose_is_rejected_even_with_a_valid_signature():
    forged = jwt.encode(
        {"profile_id": "x", "row_number": 1, "fields": {}, "has_duplicate_advisory": False, "exp": 9999999999},
        config.SESSION_SECRET, algorithm="HS256",
    )
    with pytest.raises(pis.RowTokenError):
        pis.verify_row_token(forged)


def test_a_real_session_shaped_jwt_is_rejected_by_the_row_token_verifier():
    from webapp.services import auth_service
    session = auth_service.create_session(
        {"email": "u@danwor.com", "name": "U", "first_name": "U", "family_name": ""}
    )
    with pytest.raises(pis.RowTokenError):
        pis.verify_row_token(session)


def test_a_row_token_cannot_authenticate_as_a_hub_session():
    from webapp.services import auth_service
    data = _workbook_bytes(_ALL_HEADERS, [_row()])
    preview = pis.build_preview(pis.parse_workbook(data))
    decoded = auth_service.decode_session(preview[0].token)
    # La firma verifica (mismo secreto), pero el token de fila no tiene
    # `sub` -- así que no puede autenticar a NINGÚN usuario real.
    assert decoded is None or not decoded.get("sub")


def test_expired_row_token_is_rejected_as_invalid_not_source_error(monkeypatch):
    monkeypatch.setattr(pis, "_ROW_TOKEN_MAX_AGE_S", -10)
    data = _workbook_bytes(_ALL_HEADERS, [_row()])
    preview = pis.build_preview(pis.parse_workbook(data))
    token = preview[0].token
    results = pis.confirm_rows([token], [token])
    assert results[0].status == "INVALID"  # nunca SOURCE_ERROR
    assert profile_repository.list() == []


# ============================================================================
# Phase 3.8.1 §5/§6 — preflight del ZIP del XLSX (anti amplificación)
# ============================================================================

def test_normal_small_xlsx_passes_preflight():
    data = _workbook_bytes(_ALL_HEADERS, [_row()])
    pis._preflight_xlsx_zip(data)  # no levanta


def test_non_zip_content_is_rejected_safely():
    with pytest.raises(pis.ImportFileError):
        pis._preflight_xlsx_zip(b"definitivamente no es un zip")


def test_uncompressed_content_over_the_bound_is_rejected_before_openpyxl(monkeypatch):
    monkeypatch.setattr(pis, "MAX_XLSX_UNCOMPRESSED_BYTES", 500)
    # Un XLSX real normal descomprime a bastante más de 500 bytes.
    data = _workbook_bytes(_ALL_HEADERS, [_row()])
    with pytest.raises(pis.ImportFileError) as exc_info:
        pis.parse_workbook(data)
    assert "descomprimido" in str(exc_info.value)


def test_too_many_zip_entries_is_rejected(monkeypatch):
    monkeypatch.setattr(pis, "MAX_XLSX_ENTRIES", 2)
    data = _workbook_bytes(_ALL_HEADERS, [_row()])  # un XLSX tiene > 2 entradas internas
    with pytest.raises(pis.ImportFileError):
        pis.parse_workbook(data)


def test_a_synthetic_zip_bomb_within_the_upload_size_is_rejected(monkeypatch):
    monkeypatch.setattr(pis, "MAX_XLSX_UNCOMPRESSED_BYTES", 1024)
    # ZIP chico en bytes comprimidos, pero con una entrada que declara
    # descomprimir a mucho más -- construido con contenido altamente
    # comprimible (repetición), nunca un artefacto gigante en disco.
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("[Content_Types].xml", b"0" * (64 * 1024))  # 64 KiB de ceros -> comprime a casi nada
    payload = buf.getvalue()
    assert len(payload) < 1024  # cabe holgado en cualquier MAX_UPLOAD_BYTES
    with pytest.raises(pis.ImportFileError):
        pis.parse_workbook(payload)


def test_preflight_never_extracts_files_to_disk(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "BASE_DIR", tmp_path)
    data = _workbook_bytes(_ALL_HEADERS, [_row()])
    pis.parse_workbook(data)
    assert list(tmp_path.rglob("*.xml")) == []
    assert list(tmp_path.rglob("*.xlsx")) == []


# ============================================================================
# Phase 3.8.1 §10-§13 — drift de duplicados entre preview y confirm
# ============================================================================

def test_no_drift_row_created_normally():
    data = _workbook_bytes(_ALL_HEADERS, [_row(email="juan@x.com")])
    preview = pis.build_preview(pis.parse_workbook(data))
    token = preview[0].token
    result = pis.confirm_rows([token], [token])[0]
    assert result.status == "CREATED"


def test_matching_profile_created_after_preview_requires_confirmation():
    data = _workbook_bytes(_ALL_HEADERS, [_row(email="juan@x.com")])
    preview = pis.build_preview(pis.parse_workbook(data))  # sin advisory: no existía nada
    token = preview[0].token
    assert preview[0].status == "eligible"

    # DRIFT: alguien más crea un Profile que matchea, DESPUÉS del preview.
    human_capital_service.create_profile({"first_name": "Juan", "last_name": "Perez", "email": "juan@x.com"})

    result = pis.confirm_rows([token], [token])[0]
    assert result.status == "DUPLICATE_REQUIRES_CONFIRMATION"
    assert result.token is not None
    # ZERO create de esta fila.
    assert len([p for p in profile_repository.list() if p["profile_id"] == preview[0].profile_id]) == 0


def test_confirming_the_refreshed_token_creates_the_profile_anyway():
    data = _workbook_bytes(_ALL_HEADERS, [_row(email="juan@x.com")])
    preview = pis.build_preview(pis.parse_workbook(data))
    token = preview[0].token
    human_capital_service.create_profile({"first_name": "Juan", "last_name": "Perez", "email": "juan@x.com"})

    first = pis.confirm_rows([token], [token])[0]
    assert first.status == "DUPLICATE_REQUIRES_CONFIRMATION"

    second = pis.confirm_rows([first.token], [first.token])[0]
    assert second.status == "CREATED"
    assert len([p for p in profile_repository.list() if p["profile_id"] == preview[0].profile_id]) == 1


def test_existing_advisory_explicitly_confirmed_still_creates_without_a_recheck_block():
    human_capital_service.create_profile({"first_name": "Juan", "last_name": "Perez", "email": "juan@x.com"})
    data = _workbook_bytes(_ALL_HEADERS, [_row(email="juan@x.com")])
    preview = pis.build_preview(pis.parse_workbook(data))
    assert preview[0].status == "existing_duplicate"  # advisory ya en preview
    token = preview[0].token
    result = pis.confirm_rows([token], [token])[0]
    assert result.status == "CREATED"  # el usuario ya overrideó, procede


def test_idempotent_retry_is_not_flagged_as_drift():
    # Una fila ya creada en un confirm anterior: al reintentar el MISMO
    # token, el rechequeo encuentra su PROPIO Profile en current_profiles
    # y NO debe marcarlo como DUPLICATE_REQUIRES_CONFIRMATION.
    data = _workbook_bytes(_ALL_HEADERS, [_row(email="juan@x.com")])
    preview = pis.build_preview(pis.parse_workbook(data))
    token = preview[0].token
    assert pis.confirm_rows([token], [token])[0].status == "CREATED"
    second = pis.confirm_rows([token], [token])[0]
    assert second.status == "ALREADY_COMMITTED"


def test_confirm_calls_list_profiles_exactly_once_per_batch(monkeypatch):
    calls = []
    original = human_capital_service.list_profiles
    monkeypatch.setattr(human_capital_service, "list_profiles", lambda: (calls.append(1), original())[1])
    data = _workbook_bytes(_ALL_HEADERS, [
        _row(email="a@x.com"), _row(email="b@x.com"), _row(email="c@x.com"),
    ])
    preview = pis.build_preview(pis.parse_workbook(data))
    calls.clear()  # ignora la llamada del preview
    tokens = [p.token for p in preview]
    pis.confirm_rows(tokens, tokens)
    assert len(calls) == 1