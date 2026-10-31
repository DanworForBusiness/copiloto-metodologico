"""test_datasphere_owner_profiles.py — selector "Cliente" (SOLO APP_OWNER) para Assessment ad-hoc y Publicar modelo.

Contrato:
  * el SERVIDOR decide: solo el dueño (``access.is_owner``, el mismo predicado de siempre) y solo con ``SHAREPOINT_DSP_OAUTH_URL`` configurada
    ve el selector; los demás no lo ven, no pueden llamar al endpoint (403) y SharePoint no se consulta ni al renderizar;
  * la fuente es el libro Excel privado de esa variable, leído con la integración existente (``fetch_shared_file``); el esquema se detecta por
    NOMBRE de columna (``Cliente | URL Host | URL Token | Client ID | Secret``), sin inventar columnas;
  * el navegador SOLO recibe metadatos seguros; el Client Secret se resuelve en el servidor a partir del ``profile_id`` opaco, en CADA
    ejecución (Assessment y Publicación), y no se registra ni se guarda;
  * SharePoint caído / sin token / libro inválido → estado estable y el formulario manual sigue igual.
Fixtures SINTÉTICAS: un libro .xlsx generado en tmp y un doble de Graph (sin red).
"""
from __future__ import annotations

import json
import logging
import re
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

import config
import webapp.main as main
from webapp import access
from webapp.routers import datasphere_assessment as asm_router
from webapp.routers import datasphere_publication as pub_router
from webapp.services import auth_service
from webapp.services import datasphere_assessment_service as asd
from webapp.services import datasphere_profiles as dp
from webapp.services import datasphere_publication_service as pub

REPO = Path(__file__).resolve().parents[2]
STEM = "04_Modelado_y_Desarrollo"
MODULE_URL = f"/modules/{STEM}?pack=datasphere_core"
PROFILES_URL = f"/modules/{STEM}/datasphere-profiles"
RUN_A = f"/modules/{STEM}/datasphere-assessment/run"
RUN_P = f"/modules/{STEM}/datasphere-publication/run"
ORIGIN = {"Origin": "http://testserver"}
SECRET_ACME = "sec-ACME-0123456789-zzz"
SECRET_BETA = "sec-BETA-9876543210-yyy"
SOURCE = "https://sharepoint-fixture.invalid/share/" + "x" * 8        # sentinela de test; la URL real NUNCA está en el código


def _client(email):
    c = TestClient(main.app)
    if email:
        c.cookies.set(auth_service.COOKIE_NAME, auth_service.create_session({"email": email, "name": "T", "first_name": "T", "family_name": "U"}))
    return c


def _workbook(path: Path, *, header=("Cliente", "URL Host", "URL Token", "Client ID", "Secret"), rows=None, top_meta=True) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "OAuth"
    if top_meta:
        ws.append(["Name", "Copiloto Danwor"])
        ws.append(["User ID", "usuario-tecnico"])
        ws.append([])
    ws.append(list(header))
    for r in (rows if rows is not None else [
        ("Acme", "https://acme.eu10.hcs.cloud.sap/", "https://acme.authentication.eu10.hana.ondemand.com/oauth/token", "sb-acme-id", SECRET_ACME),
        ("Beta", "https://beta.us20.hcs.cloud.sap/", "https://beta.authentication.us20.hana.ondemand.com/oauth/token", "sb-beta-id", SECRET_BETA),
    ]):
        ws.append(list(r))
    wb.save(path)
    return path


@pytest.fixture()
def world(tmp_path, monkeypatch):
    monkeypatch.setattr(access, "APP_OWNER", "owner@danwor.com")
    monkeypatch.setattr(config, "SHAREPOINT_DSP_OAUTH_URL", SOURCE)
    monkeypatch.setattr(config, "ASSESSMENT_CACHE_DIR", tmp_path / "adhoc-cache")      # el contexto de formulario por sesión no puede salir de tmp
    xlsx = _workbook(tmp_path / "perfiles.xlsx")
    calls = {"download": 0, "tokens": []}

    def fake_download(token, dest_dir):
        calls["download"] += 1
        calls["tokens"].append(token)
        target = Path(dest_dir) / "perfiles.xlsx"
        target.write_bytes(xlsx.read_bytes())
        return target
    monkeypatch.setattr(dp, "_download_workbook", fake_download)
    calls["xlsx"] = xlsx
    return calls


def _list(client, token="tok-graph", **kw):
    return client.post(PROFILES_URL, data={"graph_token": token}, headers=ORIGIN, **kw)


# ── A / B. quién ve el selector (decide el servidor) ────────────────────────
def test_app_owner_sees_the_client_selector_in_both_tools(world):
    html = _client("owner@danwor.com").get(MODULE_URL).text
    assert html.count("data-dw-profile-select") == 2
    for p in ("dsa", "dsp"):
        assert f'id="{p}-profile"' in html and f'for="{p}-profile"' in html
    assert html.count("Elegí un cliente...") == 2 and f'data-profiles-url="{PROFILES_URL}"' in html
    assert world["download"] == 0                                         # renderizar NO consulta SharePoint (carga diferida por el navegador)


def test_normal_user_does_not_see_the_selector_nor_the_forms(world):
    html = _client("user@danwor.com").get(MODULE_URL).text
    assert "data-dw-profile-select" not in html and "data-dw-owner-profile" not in html and "profiles-url" not in html
    assert world["download"] == 0


def test_owner_without_the_source_configured_gets_the_manual_form_unchanged(world, monkeypatch):
    monkeypatch.setattr(config, "SHAREPOINT_DSP_OAUTH_URL", "")
    html = _client("owner@danwor.com").get(MODULE_URL).text
    assert "data-dw-owner-profile" not in html and 'id="dsa-host"' in html and 'id="dsp-client_secret"' in html   # el formulario manual intacto
    assert _list(_client("owner@danwor.com")).json() == {"ok": False, "reason": dp.NOT_CONFIGURED, "message": dp.WARNINGS[dp.NOT_CONFIGURED]}


# ── C / D. el endpoint es solo del dueño ────────────────────────────────────
def test_non_owner_cannot_call_the_profiles_endpoint_and_sharepoint_is_never_queried(world, monkeypatch):
    monkeypatch.setattr(dp, "load_profiles", lambda t: pytest.fail("SharePoint no debe consultarse para un no-dueño"))
    assert _list(_client("user@danwor.com")).status_code == 403
    assert _list(_client(None)).status_code == 401
    r = _client("owner@danwor.com").post(PROFILES_URL, data={"graph_token": "t"})              # sin mismo origen
    assert r.status_code == 403


# ── E / F. el navegador solo recibe metadatos seguros ───────────────────────
def test_owner_response_has_only_safe_metadata_and_no_secret(world):
    r = _list(_client("owner@danwor.com"))
    data = r.json()
    assert data["ok"] is True and [p["client"] for p in data["profiles"]] == ["Acme", "Beta"]
    assert all(set(p) == {"profile_id", "client", "host", "token_url", "client_id"} for p in data["profiles"])
    assert SECRET_ACME not in r.text and SECRET_BETA not in r.text and "secret" not in r.text.lower()
    assert r.headers["cache-control"] == "no-store" and world["tokens"] == ["tok-graph"]
    ids = [p["profile_id"] for p in data["profiles"]]
    assert len(set(ids)) == 2 and all(re.fullmatch(r"[0-9a-f]{20}", i) for i in ids) and not any(s.lower() in i for i in ids for s in ("acme", "beta"))


def test_rendered_html_and_js_never_contain_a_client_secret(world):
    html = _client("owner@danwor.com").get(MODULE_URL).text
    js = (REPO / "webapp" / "static" / "datasphere_owner_profile.js").read_text(encoding="utf-8")
    for text in (html, js):
        assert SECRET_ACME not in text and SECRET_BETA not in text
    assert not re.search(r"\.secret\b|profile\.secret|client_secret\s*[:=]\s*[^\"'\s]", js)         # el JS nunca maneja un secreto
    for storage in ("localStorage", "sessionStorage"):
        assert storage not in js
    assert 'name="profile_id"' in html and 'type="password"' in html                               # el campo de secreto sigue siendo password vacío


# ── G. mapeo perfil → campos seguros (el comportamiento visual se verifica en el navegador) ───────────────
def test_the_selector_maps_profile_fields_onto_existing_form_fields():
    js = (REPO / "webapp" / "static" / "datasphere_owner_profile.js").read_text(encoding="utf-8")
    macro = (REPO / "webapp" / "templates" / "partials" / "datasphere_form_fields.html").read_text(encoding="utf-8")
    mapped = re.findall(r'setField\(form, p, "(\w+)", profile\.(\w+)\)', js)
    assert mapped == [("host", "host"), ("token_url", "token_url"), ("client_id", "client_id"), ("client", "client")]
    for field, _src in mapped:
        assert re.search(rf'field\((?:p, )?"{field}"', macro)                                      # el input existe en la macro compartida
    assert "Configurado desde SharePoint" in js and "secret.disabled = true" in js and "secret.disabled = false" in js


# ── H / I / J. el secreto se resuelve en el servidor al ejecutar ─────────────
def _capture_runs(monkeypatch):
    seen = {"asm": [], "pub": [], "activity": []}
    ok = asd.Outcome("FAILED", "SENTINEL", "msg")
    monkeypatch.setattr(asd, "run", lambda values, secret, **k: (seen["asm"].append((values, secret)), ok)[1])
    monkeypatch.setattr(pub, "run", lambda values, secret, **k: (seen["pub"].append((values, secret)), ok)[1])
    for router in (asm_router, pub_router):
        monkeypatch.setattr(router, "log_activity", lambda *a, **k: seen["activity"].append(k))
    return seen


def _form(**over):
    base = {"pack": "datasphere_core", "host": "https://acme.eu10.hcs.cloud.sap", "token_url": "https://acme.authentication.eu10.hana.ondemand.com/oauth/token",
            "datasphere_client_id": "sb-acme-id", "space": "S", "datasphere_root_object": "4AM_ACME", "client": "Acme"}
    base.update(over)
    return base


def _events(r):
    return [json.loads(line) for line in r.text.splitlines() if line.strip()]


@pytest.mark.parametrize("url, bucket", [(RUN_A, "asm"), (RUN_P, "pub")])
def test_execution_resolves_the_secret_server_side_from_the_profile_id(world, monkeypatch, url, bucket):
    seen = _capture_runs(monkeypatch)
    pid = _list(_client("owner@danwor.com")).json()["profiles"][0]["profile_id"]
    r = _client("owner@danwor.com").post(url, data=_form(profile_id=pid, graph_token="tok-run", datasphere_client_secret="atacante"), headers=ORIGIN)
    assert r.status_code == 200 and _events(r)[-1]["state"] == "FAILED"                            # el doble de run devuelve un resultado cualquiera
    ((values, secret),) = seen[bucket]
    assert secret == SECRET_ACME                                                                   # resuelto en el servidor; el valor posteado se ignora
    assert world["tokens"][-1] == "tok-run" and SECRET_ACME not in r.text                           # y nunca vuelve al navegador
    assert not hasattr(values, "client_secret") and SECRET_ACME not in json.dumps(seen["activity"], default=str)


@pytest.mark.parametrize("url, bucket", [(RUN_A, "asm"), (RUN_P, "pub")])
def test_unresolvable_profile_fails_cleanly_without_running(world, monkeypatch, url, bucket):
    seen = _capture_runs(monkeypatch)
    r = _client("owner@danwor.com").post(url, data=_form(profile_id="0" * 20, graph_token="tok"), headers=ORIGIN)
    done = _events(r)[-1]
    assert done["state"] == "FAILED" and done["error_code"] == "PROFILE_UNAVAILABLE" and done["title"] == dp.RUN_FAILURE_TITLE
    assert seen[bucket] == []                                                                       # no se ejecutó nada y no hubo secreto vacío


@pytest.mark.parametrize("url, bucket", [(RUN_A, "asm"), (RUN_P, "pub")])
def test_manual_mode_is_unchanged_without_a_profile(world, monkeypatch, url, bucket):
    seen = _capture_runs(monkeypatch)
    r = _client("owner@danwor.com").post(url, data=_form(datasphere_client_secret="manual-secret"), headers=ORIGIN)
    assert r.status_code == 200
    assert seen[bucket][0][1] == "manual-secret" and world["download"] == 0                         # sin perfil: ni toca SharePoint


def test_a_non_owner_cannot_use_a_profile_id_to_run(world, monkeypatch):
    seen = _capture_runs(monkeypatch)
    pid = _list(_client("owner@danwor.com")).json()["profiles"][0]["profile_id"]
    for url in (RUN_A, RUN_P):
        assert _client("user@danwor.com").post(url, data=_form(profile_id=pid, graph_token="t"), headers=ORIGIN).status_code == 403
    assert seen["asm"] == [] and seen["pub"] == []


# ── K. SharePoint falla: el formulario manual sigue ─────────────────────────
@pytest.mark.parametrize("token, mutate, reason", [
    ("", None, dp.NO_TOKEN),
    ("t", "boom", dp.UNAVAILABLE),
    ("t", "badheader", dp.INVALID_WORKBOOK),
    ("t", "empty", dp.EMPTY),
])
def test_sharepoint_failures_degrade_to_a_stable_state(world, monkeypatch, tmp_path, token, mutate, reason):
    if mutate == "boom":
        monkeypatch.setattr(dp, "_download_workbook", lambda t, d: (_ for _ in ()).throw(RuntimeError("grafo caído: detalle interno")))
    elif mutate == "badheader":
        bad = _workbook(tmp_path / "bad.xlsx", header=("Cliente", "URL Host", "URL Token", "Client ID", "Otra"))
        monkeypatch.setattr(dp, "_download_workbook", lambda t, d: bad)
    elif mutate == "empty":
        empty = _workbook(tmp_path / "empty.xlsx", rows=[])
        monkeypatch.setattr(dp, "_download_workbook", lambda t, d: empty)
    r = _list(_client("owner@danwor.com"), token=token)
    assert r.status_code == 200 and r.json() == {"ok": False, "reason": reason, "message": dp.WARNINGS[reason]}
    assert "detalle interno" not in r.text                                                          # nada del error de SharePoint llega al navegador
    html = _client("owner@danwor.com").get(MODULE_URL).text
    assert 'id="dsa-host"' in html and 'id="dsp-host"' in html                                      # la página y los formularios no dependen de la fuente


# ── esquema del libro / identidad opaca ─────────────────────────────────────
def test_workbook_is_parsed_by_header_name_not_position(tmp_path):
    path = _workbook(tmp_path / "w.xlsx", header=("Secret", "Client ID", "URL Token", "URL Host", "Cliente"),
                     rows=[(SECRET_ACME, "id1", "https://t/oauth/token", "https://h/", "Acme")])
    (p,) = dp.parse_workbook(path)
    assert (p.client, p.host, p.token_url, p.client_id, p.secret) == ("Acme", "https://h/", "https://t/oauth/token", "id1", SECRET_ACME)


def test_incomplete_rows_are_skipped_and_duplicates_keep_the_first(tmp_path):
    rows = [("Acme", "h1", "t1", "i1", "s1"), ("", "h", "t", "i", "s"), ("Beta", "h2", "", "i2", "s2"), ("acme", "h3", "t3", "i3", "s3"), ("Gamma", "h4", "t4", "i4", "s4")]
    got = dp.parse_workbook(_workbook(tmp_path / "w.xlsx", rows=rows))
    assert [(p.client, p.host) for p in got] == [("Acme", "h1"), ("Gamma", "h4")]


def test_profile_repr_and_public_view_hide_the_secret(tmp_path):
    (p, *_) = dp.parse_workbook(_workbook(tmp_path / "w.xlsx"))
    assert SECRET_ACME not in repr(p) and SECRET_ACME not in json.dumps(p.public()) and p.profile_id == dp.profile_id_for("  ACME ")


def test_the_workbook_is_downloaded_with_the_existing_integration_into_a_temp_dir(monkeypatch, tmp_path):
    """`_download_workbook` reusa `generate_catalog.fetch_shared_file` (Graph /shares + token delegado) y deja el libro SOLO en un temporal."""
    import sys
    monkeypatch.setattr(config, "SHAREPOINT_DSP_OAUTH_URL", SOURCE)
    sys.path.insert(0, str(REPO / "tools" / "catalog"))
    import generate_catalog as gen
    seen = {}
    xlsx = _workbook(tmp_path / "src.xlsx")

    def fake_fetch(url, token, dest, **kw):
        seen.update(url=url, token=token, dest=Path(dest), kw=kw)
        (Path(dest) / "x.xlsx").write_bytes(xlsx.read_bytes())
        return "ok"
    monkeypatch.setattr(gen, "fetch_shared_file", fake_fetch)
    status, profiles = dp.load_profiles("tok-delegado")
    assert status == dp.OK and len(profiles) == 2
    assert seen["url"] == SOURCE and seen["token"] == "tok-delegado" and "cache_root" not in seen["kw"]
    assert not seen["dest"].exists()                                                              # el temporal con secretos ya no existe


def test_no_url_is_hardcoded_and_no_new_sharepoint_client_exists():
    for rel in ("webapp/services/datasphere_profiles.py", "webapp/static/datasphere_owner_profile.js",
                "webapp/templates/partials/datasphere_form_fields.html", "webapp/routers/datasphere_assessment.py"):
        text = (REPO / rel).read_text(encoding="utf-8")
        assert not re.search(r"sharepoint\.com|:x:/s/", text), rel
    src = (REPO / "webapp/services/datasphere_profiles.py").read_text(encoding="utf-8")
    assert "fetch_shared_file" in src and "msal" not in src.lower() and "requests" not in src and "urllib" not in src
    assert 'SHAREPOINT_DSP_OAUTH_URL = _env("SHAREPOINT_DSP_OAUTH_URL", "")' in (REPO / "config.py").read_text(encoding="utf-8")   # única fuente, sin valor


# ── L. nada de credenciales en logs ─────────────────────────────────────────
def test_no_credential_reaches_the_logs(world, monkeypatch, caplog):
    caplog.set_level(logging.DEBUG)
    seen = _capture_runs(monkeypatch)
    c = _client("owner@danwor.com")
    pid = _list(c, token="tok-secreto-graph").json()["profiles"][1]["profile_id"]
    c.post(RUN_A, data=_form(profile_id=pid, graph_token="tok-secreto-graph"), headers=ORIGIN)
    monkeypatch.setattr(dp, "_download_workbook", lambda t, d: (_ for _ in ()).throw(RuntimeError(f"fallo con {SECRET_BETA}")))
    _list(c, token="tok-secreto-graph")
    text = caplog.text
    for s in (SECRET_ACME, SECRET_BETA, "tok-secreto-graph", SOURCE):
        assert s not in text, s
    assert SECRET_BETA not in json.dumps(seen["activity"], default=str)
