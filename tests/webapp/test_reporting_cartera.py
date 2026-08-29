"""
test_reporting_cartera.py — Contrato de datos reales de Cartera (Reporting
Integral real-data Fase 1). Determinístico, sin Graph real: monkeypatchea
`core.microsoft365` (mismo criterio que los tests de Staffing) y construye un
XLSX real con la tabla `Tabla24` en memoria.
"""
from __future__ import annotations

import io

import openpyxl
import pytest
from openpyxl.worksheet.table import Table

from webapp.services import reporting_repository as repo_mod
from webapp.services.reporting_repository import (
    ReportingConfigurationError,
    ReportingGraphError,
    ReportingRepository,
    ReportingSourceError,
)
from webapp.services.reporting_service import ReportingService

_SITE, _HP, _PROY, _CLI = "site-1", "hp-list", "proy-list", "cli-list"


def _items(rows):
    return [{"id": r.get("id", ""), "fields": r} for r in rows]


def _facturadas_xlsx(pairs, *, headers=("Proyecto Base", "Suma de Cantidad"),
                     lead_rows=0, decoy_sheet=False):
    """Construye el XLSX de Facturadas. `lead_rows`/`decoy_sheet` permiten
    ubicar la tabla FUERA de A1 / en una hoja posterior, para probar que la
    extracción no asume posición (read_only no expone la metadata de tabla)."""
    wb = openpyxl.Workbook()
    if decoy_sheet:
        decoy = wb.active
        decoy.title = "Portada"
        decoy.append(["Reporte", "Horas Facturadas"])
        ws = wb.create_sheet("Facturado")
    else:
        ws = wb.active
        ws.title = "Facturado"
    for _ in range(lead_rows):
        ws.append(["", ""])  # filas de relleno antes del header
    header_row = len(ws["A"]) + 1 if lead_rows else 1
    ws.append(list(headers))
    for proyecto_base, cantidad in pairs:
        ws.append([proyecto_base, cantidad])
    ws.add_table(Table(displayName="Tabla24", ref=f"A{header_row}:B{header_row + len(pairs)}"))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


_CP = "cp-list"


def _patch(monkeypatch, *, horas, proyectos, clientes, facturadas, cp_columns=None, client_partners=None):
    def fake_paged(url, access_token=None, **kw):
        if "/columns" in url:                       # metadata de columnas de Cliente
            return cp_columns if cp_columns is not None else []
        if _CP in url:                              # lista target de Client Partner
            return _items(client_partners or [])
        if _HP in url:
            return _items(horas)
        if _PROY in url:
            return _items(proyectos)
        if _CLI in url:
            return _items(clientes)
        return []

    def fake_bytes(url, access_token=None, **kw):
        if isinstance(facturadas, Exception):
            raise facturadas
        return facturadas

    monkeypatch.setattr(repo_mod.microsoft365, "graph_get_paged_values", fake_paged)
    monkeypatch.setattr(repo_mod.microsoft365, "graph_get_bytes", fake_bytes)


_CP_COLUMNS = [{"name": "Title"}, {"name": "ClientPartner", "lookup": {"listId": _CP}}]


def _repo():
    return ReportingRepository(
        site_id=_SITE, horas_proyectos_list_id=_HP,
        proyectos_list_id=_PROY, cliente_list_id=_CLI,
    )


# --- Repository: normalización + joins ---------------------------------------

def test_repo_joins_lists_and_xlsx_and_normalizes(monkeypatch):
    # Nombres internos REALES (live smoke): A_x00f1_o, Status, ID_x0020_...,
    # Cliente.Title, ClientPartnerLookupId → target descubierto de /columns.
    _patch(
        monkeypatch,
        horas=[{"ProyectoLookupId": "p1", "EmpresaLookupId": "c1",
                "A_x00f1_o": "2026", "Mes": "Marzo", "HsTrabajadas": "100", "HsVendidas": "120", "HsFacturadas": "110"}],
        proyectos=[{"id": "p1", "Title": "Implementación SADI v2", "ID_x0020_Proyecto_x0020_Danwor": "DAN-1042",
                    "TipodeProyecto": "Proyecto", "FechadeInicio": "2026-01-15", "Status": "En curso"}],
        clientes=[{"id": "c1", "Title": "YPF", "ClientPartnerLookupId": "pm1"}],
        cp_columns=_CP_COLUMNS,
        client_partners=[{"id": "pm1", "Title": "Diego Ramos"}],
        facturadas=_facturadas_xlsx([("Implementación SADI v2", 290)]),
    )
    facts = _repo().get_cartera_facts("tok")
    assert len(facts) == 1
    f = facts[0]
    assert f["cliente"] == "YPF" and f["clientPartner"] == "Diego Ramos"
    assert f["proyecto"] == "Implementación SADI v2" and f["idProyecto"] == "DAN-1042"
    assert f["tipoProyecto"] == "Proyecto" and f["estado"] == "En curso" and f["anioInicio"] == "2026"
    assert f["hsTrabajadas"] == 100 and f["hsVendidas"] == 120 and f["hsCertificadas"] == 110
    assert f["facturadasProyecto"] == 290  # join XLSX por nombre normalizado


def test_graph_form_lookup_ids_and_title_from_joined_project(monkeypatch):
    # Reproduce el bug del live-smoke (resultado vacío): Graph entrega los
    # lookups como ProyectoLookupId/EmpresaLookupId (no ProyectoId/EmpresaId) y
    # NO trae el título en la fila de Horas — el join debe resolver igual.
    _patch(
        monkeypatch,
        horas=[{"ProyectoLookupId": "7", "EmpresaLookupId": "3", "A_x00f1_o": "2026", "Mes": "Marzo",
                "HsTrabajadas": "100", "HsVendidas": "120", "HsFacturadas": "110"}],
        proyectos=[{"id": "7", "Title": "Implementación SADI v2", "ID_x0020_Proyecto_x0020_Danwor": "DAN-1042",
                    "TipodeProyecto": "Proyecto", "FechadeInicio": "2026-01-15", "Status": "En curso"}],
        clientes=[{"id": "3", "Title": "YPF", "ClientPartnerLookupId": "pm1"}],
        cp_columns=_CP_COLUMNS,
        client_partners=[{"id": "pm1", "Title": "Diego Ramos"}],
        facturadas=_facturadas_xlsx([("Implementación SADI v2", 290)]),
    )
    facts = _repo().get_cartera_facts("tok")
    assert len(facts) == 1
    f = facts[0]
    assert f["proyecto"] == "Implementación SADI v2" and f["cliente"] == "YPF"
    assert f["clientPartner"] == "Diego Ramos" and f["idProyecto"] == "DAN-1042"
    assert f["tipoProyecto"] == "Proyecto" and f["facturadasProyecto"] == 290


def test_tiempo_y_materiales_maps_to_mantenimiento(monkeypatch):
    _patch(
        monkeypatch,
        horas=[{"ProyectoId": "p1", "EmpresaId": "c1", "Proyecto": "Mant BI",
                "Año": "2026", "Mes": "Junio", "HsTrabajadas": "60", "HsVendidas": "60", "HsFacturadas": "60"}],
        proyectos=[{"id": "p1", "TipodeProyecto": "Tiempo & Materiales", "Estado": "En curso"}],
        clientes=[{"id": "c1", "Cliente": "Molinos", "Client Partner": "María López"}],
        facturadas=_facturadas_xlsx([("Otro", 0)]),
    )
    assert _repo().get_cartera_facts("tok")[0]["tipoProyecto"] == "Mantenimiento"


def test_facturadas_zero_when_project_absent_in_xlsx(monkeypatch):
    _patch(
        monkeypatch,
        horas=[{"ProyectoId": "p1", "EmpresaId": "c1", "Proyecto": "SADI",
                "Año": "2026", "Mes": "1", "HsTrabajadas": "10", "HsVendidas": "10", "HsFacturadas": "10"}],
        proyectos=[{"id": "p1", "TipodeProyecto": "Proyecto", "Estado": "En curso"}],
        clientes=[{"id": "c1", "Cliente": "YPF", "Client Partner": "Diego Ramos"}],
        facturadas=_facturadas_xlsx([("Proyecto Distinto", 999)]),
    )
    assert _repo().get_cartera_facts("tok")[0]["facturadasProyecto"] == 0.0


def test_missing_referenced_project_or_client_yields_empty_dims(monkeypatch):
    _patch(
        monkeypatch,
        horas=[{"ProyectoId": "MISSING", "EmpresaId": "MISSING", "Proyecto": "Huérfano",
                "Año": "2026", "Mes": "5", "HsTrabajadas": "5", "HsVendidas": "5", "HsFacturadas": "5"}],
        proyectos=[{"id": "p1", "TipodeProyecto": "Proyecto"}],
        clientes=[{"id": "c1", "Cliente": "YPF"}],
        facturadas=_facturadas_xlsx([("x", 0)]),
    )
    f = _repo().get_cartera_facts("tok")[0]
    assert f["cliente"] == "" and f["clientPartner"] == "" and f["tipoProyecto"] == ""
    assert f["proyecto"] == "Huérfano"  # el título viene de la fila mensual


def test_repo_consumes_all_paged_items(monkeypatch):
    _patch(
        monkeypatch,
        horas=[{"ProyectoId": "p1", "EmpresaId": "c1", "Proyecto": "A", "Año": "2026", "Mes": "1",
                "HsTrabajadas": "1", "HsVendidas": "1", "HsFacturadas": "1"} for _ in range(5)],
        proyectos=[{"id": "p1", "TipodeProyecto": "Proyecto"}],
        clientes=[{"id": "c1", "Cliente": "YPF"}],
        facturadas=_facturadas_xlsx([("A", 3)]),
    )
    assert len(_repo().get_cartera_facts("tok")) == 5


def test_facturadas_table_found_when_not_at_a1_and_not_first_sheet(monkeypatch):
    # Regresión del live-smoke hardening: la tabla se ancla al header real, no
    # a A1/primera hoja (read_only no expone ws.tables). Aún así debe unir.
    _patch(
        monkeypatch,
        horas=[{"ProyectoId": "p1", "EmpresaId": "c1", "Proyecto": "SADI",
                "Año": "2026", "Mes": "1", "HsTrabajadas": "10", "HsVendidas": "10", "HsFacturadas": "10"}],
        proyectos=[{"id": "p1", "TipodeProyecto": "Proyecto", "Estado": "En curso"}],
        clientes=[{"id": "c1", "Cliente": "YPF", "Client Partner": "Diego Ramos"}],
        facturadas=_facturadas_xlsx([("SADI", 290)], lead_rows=3, decoy_sheet=True),
    )
    assert _repo().get_cartera_facts("tok")[0]["facturadasProyecto"] == 290


def test_facturadas_table_missing_header_raises_source_error(monkeypatch):
    bad = _facturadas_xlsx([("SADI", 1)], headers=("Otra Cosa", "Valor"))
    _patch(monkeypatch, horas=[], proyectos=[], clientes=[], facturadas=bad)
    with pytest.raises(ReportingSourceError) as exc:
        _repo().get_cartera_facts("tok")
    assert exc.value.stage == "table"


def test_malformed_workbook_raises_source_error(monkeypatch):
    _patch(
        monkeypatch,
        horas=[], proyectos=[], clientes=[],
        facturadas=b"esto no es un xlsx",
    )
    with pytest.raises(ReportingSourceError) as exc:
        _repo().get_cartera_facts("tok")
    assert exc.value.stage == "workbook"


def test_missing_site_id_raises_configuration_error():
    repo = ReportingRepository(site_id="", horas_proyectos_list_id=_HP,
                               proyectos_list_id=_PROY, cliente_list_id=_CLI)
    with pytest.raises(ReportingConfigurationError):
        repo.get_cartera_facts("tok")


def test_graph_http_error_wrapped(monkeypatch):
    import urllib.error

    def boom(url, access_token=None, **kw):
        raise urllib.error.HTTPError(url, 403, "boom", None, None)

    monkeypatch.setattr(repo_mod.microsoft365, "graph_get_paged_values", boom)
    with pytest.raises(ReportingGraphError) as exc:
        _repo().get_cartera_facts("tok")
    assert exc.value.status_code == 403


def test_client_partner_resolved_from_discovered_lookup_target(monkeypatch):
    # El target del lookup CP se DESCUBRE de /columns (lookup.listId), no se
    # asume; se carga una vez y se mapea id→Title.
    _patch(
        monkeypatch,
        horas=[{"ProyectoLookupId": "p1", "EmpresaLookupId": "c1", "A_x00f1_o": "2026", "Mes": "1",
                "HsTrabajadas": "1", "HsVendidas": "1", "HsFacturadas": "1"}],
        proyectos=[{"id": "p1", "TipodeProyecto": "Proyecto"}],
        clientes=[{"id": "c1", "Title": "YPF", "ClientPartnerLookupId": "pm9"}],
        cp_columns=_CP_COLUMNS,
        client_partners=[{"id": "pm9", "Title": "Ana Gómez"}],
        facturadas=_facturadas_xlsx([("x", 0)]),
    )
    # Protege el descubrimiento del target de CP (lookup.listId de /columns).
    assert _repo().get_cartera_facts("tok")[0]["clientPartner"] == "Ana Gómez"


def test_client_partner_blank_when_target_not_discoverable_no_guess(monkeypatch):
    # Sin columna lookup ClientPartner en /columns → NO se asume target → CP en
    # blanco (nunca se fabrica).
    _patch(
        monkeypatch,
        horas=[{"ProyectoLookupId": "p1", "EmpresaLookupId": "c1", "A_x00f1_o": "2026", "Mes": "1",
                "HsTrabajadas": "1", "HsVendidas": "1", "HsFacturadas": "1"}],
        proyectos=[{"id": "p1", "TipodeProyecto": "Proyecto"}],
        clientes=[{"id": "c1", "Title": "YPF", "ClientPartnerLookupId": "pm9"}],
        cp_columns=[{"name": "Title"}],  # sin lookup ClientPartner
        client_partners=[{"id": "pm9", "Title": "Ana Gómez"}],
        facturadas=_facturadas_xlsx([("x", 0)]),
    )
    assert _repo().get_cartera_facts("tok")[0]["clientPartner"] == ""


# ============================================================================
# P5 — Mantenimiento nunca dispara la XLSX de Facturadas (owner live QA de
# performance, Cartera ~17s)
# ============================================================================

def test_include_facturadas_false_never_downloads_xlsx(monkeypatch):
    calls = {"bytes": 0}

    def counting_bytes(url, access_token=None, **kw):
        calls["bytes"] += 1
        raise AssertionError("no debería descargarse la XLSX con include_facturadas=False")

    _patch(
        monkeypatch,
        horas=[{"ProyectoLookupId": "p1", "EmpresaLookupId": "c1", "A_x00f1_o": "2026", "Mes": "Agosto",
                "HsTrabajadas": "10", "HsVendidas": "12", "HsFacturadas": "11"}],
        proyectos=[{"id": "p1", "Title": "Mantenimiento X", "TipodeProyecto": "Tiempo & Materiales"}],
        clientes=[{"id": "c1", "Title": "Andesmar"}],
        facturadas=Exception("no debería tocarse"),
    )
    monkeypatch.setattr(repo_mod.microsoft365, "graph_get_bytes", counting_bytes)
    facts = _repo().get_cartera_facts("tok", include_facturadas=False)
    assert calls["bytes"] == 0
    assert facts[0]["facturadasProyecto"] == 0.0


def test_include_facturadas_true_still_downloads_xlsx_default_preserved(monkeypatch):
    # Default sin cambios (retrocompat) — Proyecto/`get_cartera()` combinado
    # SIGUEN descargando la XLSX como siempre.
    _patch(
        monkeypatch,
        horas=[{"ProyectoLookupId": "p1", "EmpresaLookupId": "c1", "A_x00f1_o": "2026", "Mes": "Enero",
                "HsTrabajadas": "10", "HsVendidas": "12", "HsFacturadas": "11"}],
        proyectos=[{"id": "p1", "Title": "SADI v2", "TipodeProyecto": "Proyecto"}],
        clientes=[{"id": "c1", "Title": "YPF"}],
        facturadas=_facturadas_xlsx([("SADIV2", 99)]),
    )
    facts = _repo().get_cartera_facts("tok")  # sin kwarg = True
    assert facts[0]["facturadasProyecto"] == 99


# ============================================================================
# HOT WINDOW / COLD HISTORY (owner corrective — Cartera hidrataba
# ~2023→2026 completo aunque el default aterriza en el año calendario
# actual). `hot_years` intenta `$filter` OData con fallback automático a
# lectura completa + filtro Python (MISMO patrón ya probado en producción
# por `staffing_repository._fetch_asignaciones`).
# ============================================================================

_ALL_HORAS = [
    {"ProyectoLookupId": "p1", "EmpresaLookupId": "c1", "A_x00f1_o": "2026", "Mes": "Marzo",
     "HsTrabajadas": "10", "HsVendidas": "12", "HsFacturadas": "11"},
    {"ProyectoLookupId": "p1", "EmpresaLookupId": "c1", "A_x00f1_o": "2025", "Mes": "Marzo",
     "HsTrabajadas": "20", "HsVendidas": "22", "HsFacturadas": "21"},
    {"ProyectoLookupId": "p1", "EmpresaLookupId": "c1", "A_x00f1_o": "2023", "Mes": "Marzo",
     "HsTrabajadas": "30", "HsVendidas": "32", "HsFacturadas": "31"},
]
_PROYECTOS_HOT = [{"id": "p1", "Title": "Mantenimiento X", "TipodeProyecto": "Tiempo & Materiales", "Status": "En curso"}]
_CLIENTES_HOT = [{"id": "c1", "Title": "Andesmar"}]


def test_hot_years_filter_accepted_scopes_the_graph_read(monkeypatch):
    def fake_paged(url, access_token=None, **kw):
        if "/columns" in url:
            return []
        if _HP in url:
            assert "$filter" in url  # el filtro SÍ viajó en la URL
            assert kw.get("extra_headers", {}).get("Prefer") == "HonorNonIndexedQueriesWarningMayFailRandomly"
            return _items([h for h in _ALL_HORAS if h["A_x00f1_o"] in ("2026", "2025")])
        if _PROY in url:
            return _items(_PROYECTOS_HOT)
        if _CLI in url:
            return _items(_CLIENTES_HOT)
        return []
    monkeypatch.setattr(repo_mod.microsoft365, "graph_get_paged_values", fake_paged)
    facts = _repo().get_cartera_facts("tok", include_facturadas=False, hot_years=["2026", "2025"])
    assert {f["anio"] for f in facts} == {"2026", "2025"}


def test_hot_years_filter_rejected_falls_back_to_full_read_and_python_filter(monkeypatch):
    def fake_paged(url, access_token=None, **kw):
        if "/columns" in url:
            return []
        if _HP in url:
            if "$filter" in url:
                raise RuntimeError("Graph rechaza $filter sobre columna no indexada")
            return _items(_ALL_HORAS)  # lectura completa (fallback)
        if _PROY in url:
            return _items(_PROYECTOS_HOT)
        if _CLI in url:
            return _items(_CLIENTES_HOT)
        return []
    monkeypatch.setattr(repo_mod.microsoft365, "graph_get_paged_values", fake_paged)
    # Nunca se rompe — el fallback filtra en Python y devuelve SOLO el año pedido.
    facts = _repo().get_cartera_facts("tok", include_facturadas=False, hot_years=["2026"])
    assert {f["anio"] for f in facts} == {"2026"}


def test_hot_years_none_reads_full_history_unfiltered(monkeypatch):
    # Retrocompat exacta — Proyecto/`get_cartera()`/tests existentes NUNCA
    # ven un `$filter` que no pidieron.
    def fake_paged(url, access_token=None, **kw):
        if "/columns" in url:
            return []
        if _HP in url:
            assert "$filter" not in url
            return _items(_ALL_HORAS)
        if _PROY in url:
            return _items(_PROYECTOS_HOT)
        if _CLI in url:
            return _items(_CLIENTES_HOT)
        return []
    monkeypatch.setattr(repo_mod.microsoft365, "graph_get_paged_values", fake_paged)
    facts = _repo().get_cartera_facts("tok", include_facturadas=False)  # hot_years=None (default)
    assert {f["anio"] for f in facts} == {"2026", "2025", "2023"}


# --- P3 (field narrowing): $select en Proyectos/Cliente/Horas de Cartera --

def test_get_cartera_facts_selects_narrowed_fields_when_supported(monkeypatch):
    # Perf investigation (owner corrective — Cartera ~13s pese a la ventana
    # caliente): $select acotado a las columnas realmente leídas, NUNCA el
    # body completo del item — reduce el payload por-item de Proyectos/
    # Cliente/Horas sin filtrar ninguna fila (seguro por construcción).
    seen_urls = []

    def fake_paged(url, access_token=None, **kw):
        seen_urls.append(url)
        if "/columns" in url:
            return []
        if _HP in url:
            return _items(_ALL_HORAS)
        if _PROY in url:
            return _items(_PROYECTOS_HOT)
        if _CLI in url:
            return _items(_CLIENTES_HOT)
        return []
    monkeypatch.setattr(repo_mod.microsoft365, "graph_get_paged_values", fake_paged)
    facts = _repo().get_cartera_facts("tok", include_facturadas=False, hot_years=["2026"])
    assert facts  # sigue devolviendo datos correctos con $select activo

    proy_url = next(u for u in seen_urls if _PROY in u)
    cli_url = next(u for u in seen_urls if _CLI in u)
    hp_url = next(u for u in seen_urls if _HP in u)
    assert "$select=" in proy_url and "$select=" in cli_url and "$select=" in hp_url


def test_get_cartera_facts_falls_back_to_full_item_when_select_rejected(monkeypatch):
    # Si Graph rechaza el $select (candidato de lookup inexistente en el
    # schema real, ver CLI_CP_LOOKUP/HP_PROYECTO_LOOKUP — dos formas
    # candidatas, nunca confirmadas contra Graph real en este entorno), el
    # fetch NUNCA se rompe — cae a leer el item completo sin $select.
    def fake_paged(url, access_token=None, **kw):
        if "/columns" in url:
            return []
        if "$select" in url:
            raise RuntimeError("Graph rechaza este $select")
        if _HP in url:
            return _items(_ALL_HORAS)
        if _PROY in url:
            return _items(_PROYECTOS_HOT)
        if _CLI in url:
            return _items(_CLIENTES_HOT)
        return []
    monkeypatch.setattr(repo_mod.microsoft365, "graph_get_paged_values", fake_paged)
    facts = _repo().get_cartera_facts("tok", include_facturadas=False, hot_years=["2026"])
    assert {f["anio"] for f in facts} == {"2026"}  # el fetch entero sigue funcionando


def test_cartera_mantenimiento_anio_estado_pairs_mantenimiento_only_distinct_sorted(monkeypatch):
    _patch(
        monkeypatch,
        horas=[
            {"ProyectoLookupId": "p1", "EmpresaLookupId": "c1", "A_x00f1_o": "2026", "Mes": "Marzo"},
            {"ProyectoLookupId": "p1", "EmpresaLookupId": "c1", "A_x00f1_o": "2026", "Mes": "Abril"},  # duplicado de año
            {"ProyectoLookupId": "p1", "EmpresaLookupId": "c1", "A_x00f1_o": "2024", "Mes": "Enero"},
            {"ProyectoLookupId": "p2", "EmpresaLookupId": "c1", "A_x00f1_o": "2025", "Mes": "Enero"},  # Proyecto, no Mantenimiento
        ],
        proyectos=[
            {"id": "p1", "Title": "Mantenimiento X", "TipodeProyecto": "Tiempo & Materiales", "Status": "En curso"},
            {"id": "p2", "Title": "Proyecto Y", "TipodeProyecto": "Proyecto", "Status": "En curso"},
        ],
        clientes=_CLIENTES_HOT,
        facturadas=b"",
    )
    pairs = _repo().get_cartera_mantenimiento_anio_estado_pairs("tok")
    assert pairs == [{"anio": "2024", "estado": "En curso"}, {"anio": "2026", "estado": "En curso"}]


# --- Service: hot-window defaults + Proyecto never scoped --------------------

class _HotYearsRecordingRepo:
    def __init__(self):
        self.calls = []
        # CANONICAL SOURCE DECISION (owner 2026-08-18): Mantenimiento ahora
        # también lee Time Logs (para "Horas Trabajadas") — separado de
        # `calls` (Horas por proyectos) para no romper las aserciones
        # `calls[-1]` ya existentes.
        self.time_logs_calls = []

    def get_cartera_facts(self, access_token, *, include_facturadas=True, hot_years=None):
        self.calls.append({"include_facturadas": include_facturadas, "hot_years": hot_years})
        return []

    def get_cartera_facts_with_project_dimension(self, access_token, *, include_facturadas=True, hot_years=None, cache_scope=None):
        # UNION GRAIN corrective (owner 2026-08-18): Mantenimiento/combo
        # ahora piden la dimensión de proyecto junto con los facts —
        # MISMO tracking de `calls` (es el mismo fetch subyacente, ver
        # reporting_repository.py).
        self.calls.append({"include_facturadas": include_facturadas, "hot_years": hot_years})
        return [], {}

    def get_time_logs_worked_hours(self, access_token, *, hot_years=None, cache_scope=None):
        self.time_logs_calls.append(hot_years)
        return []

    def get_cartera_mantenimiento_anio_estado_pairs(self, access_token):
        return []


def test_service_default_hot_years_is_current_and_previous_year(monkeypatch):
    class _FixedDatetime:
        @staticmethod
        def now():
            import datetime as _dt
            return _dt.datetime(2026, 8, 15)

    import webapp.services.reporting_service as service_mod
    monkeypatch.setattr(service_mod, "datetime", _FixedDatetime)
    repo = _HotYearsRecordingRepo()
    ReportingService(repo).get_cartera_mantenimiento("tok")
    # "Horas por proyectos" (SharePoint) preserva la ventana caliente
    # especulativa de siempre (año actual + anterior) — Año/Mes puede
    # tildar el anterior sin round-trip.
    assert repo.calls[-1]["hot_years"] == ["2026", "2025"]
    # CALL BUDGET corrective (owner 2026-08-18): Time Logs es CONTEXTUAL,
    # no una ventana especulativa — sin selección explícita, SOLO el año
    # calendario actual (nunca el anterior "porque sí"). Evitar el 4x de
    # llamadas Graph reportado por el owner.
    assert repo.time_logs_calls[-1] == ["2026"]


def test_service_explicit_anios_overrides_default_hot_window():
    repo = _HotYearsRecordingRepo()
    ReportingService(repo).get_cartera_mantenimiento("tok", anios=["2024"])
    assert repo.calls[-1]["hot_years"] == ["2024"]
    assert repo.time_logs_calls[-1] == ["2024"]


def test_service_get_cartera_proyecto_never_receives_hot_years():
    repo = _HotYearsRecordingRepo()
    ReportingService(repo).get_cartera_proyecto("tok")
    assert repo.calls[-1]["hot_years"] is None


def test_service_get_cartera_combo_never_receives_hot_years(monkeypatch):
    # `get_cartera()` (retrocompat combo) sigue SIN acotar "Horas por
    # proyectos" — Proyecto necesita la vida completa, mismo criterio de
    # siempre. Y desde OPTION A (owner decision) Time Logs TAMPOCO se acota
    # acá: la mitad Proyecto ahora deriva `trabajadas` de Time Logs con grano
    # de VIDA, así que un scope contextual daría un "total de por vida" de un
    # solo año. El default contextual sigue siendo de
    # `get_cartera_mantenimiento` (ver su propio test más arriba), donde es una
    # decisión de PERFORMANCE: Mantenimiento filtra por Año/Mes de todos modos.
    class _FixedDatetime:
        @staticmethod
        def now():
            import datetime as _dt
            return _dt.datetime(2026, 8, 15)

    import webapp.services.reporting_service as service_mod
    monkeypatch.setattr(service_mod, "datetime", _FixedDatetime)
    repo = _HotYearsRecordingRepo()
    ReportingService(repo).get_cartera("tok")
    assert repo.calls[-1]["hot_years"] is None
    assert repo.time_logs_calls[-1] is None      # vida completa, no un solo año


# --- Service: derivaciones + shaping -----------------------------------------

class _FakeRepo:
    def __init__(self, facts, time_logs=None):
        self._facts = facts
        # CANONICAL SOURCE DECISION (owner 2026-08-18): Mantenimiento
        # deriva "Horas Trabajadas" de Time Logs, no de `hsTrabajadas` —
        # `time_logs=None` (default) ⇒ sin horas registradas, trabajadas=0.
        self._time_logs = time_logs if time_logs is not None else []

    def get_cartera_facts(self, access_token, *, include_facturadas=True, hot_years=None):
        return self._facts

    def get_cartera_facts_with_project_dimension(self, access_token, *, include_facturadas=True, hot_years=None, cache_scope=None):
        # UNION GRAIN corrective — estos tests NO ejercitan el camino
        # Time-Logs-only (sus fixtures ya cubren la MISMA key desde
        # "Horas por proyectos"), project_dim vacío es seguro/no-op acá.
        return self._facts, {}

    def get_time_logs_worked_hours(self, access_token, *, hot_years=None, cache_scope=None):
        return self._time_logs

    def get_cartera_mantenimiento_anio_estado_pairs(self, access_token):
        return []


def _tl_fact(**kw):
    """Fact de Time Logs mínimo para tests de servicio (grano proyecto×año×
    mes que `_trab_by_proyecto_by_year` consume) — `mes` en formato
    'YYYY-MM' (YA normalizado por el repository, ver `_time_logs_mes`)."""
    base = dict(proyecto="Mantenimiento Costeo Directo", anio="2026", mes="2026-08", horas=134.0)
    base.update(kw)
    return base


def _mant_fact(**kw):
    base = dict(tipoProyecto="Mantenimiento", cliente="Andesmar", clientPartner="María López",
                proyecto="Mantenimiento Costeo Directo", idProyecto="", estado="En curso",
                anio="2026", mes="Agosto", anioInicio="2026",
                hsTrabajadas=134.0, hsVendidas=140.0, hsCertificadas=128.0, facturadasProyecto=0.0)
    base.update(kw)
    return base


def _proy_fact(**kw):
    base = dict(tipoProyecto="Proyecto", cliente="YPF", clientPartner="Diego Ramos",
                proyecto="Implementación SADI v2", idProyecto="DAN-1042", estado="En curso",
                anio="2026", mes="Enero", anioInicio="2026",
                hsTrabajadas=200.0, hsVendidas=220.0, hsCertificadas=190.0, facturadasProyecto=290.0)
    base.update(kw)
    return base


def test_service_mantenimiento_shape_diff_and_fecha():
    # CANONICAL SOURCE DECISION (owner 2026-08-18): "trabajadas" ya no lee
    # `hsTrabajadas` (SharePoint) — se deriva de Time Logs, MISMO join
    # `_norm_key(proyecto)`×año×mes que Rentabilidad. El fixture de Time
    # Logs replica el MISMO valor (134) que antes tenía `hsTrabajadas` —
    # los números esperados (diff=-6, etc.) no cambian, solo su fuente.
    out = ReportingService(
        _FakeRepo([_mant_fact()], time_logs=[_tl_fact(horas=134.0)])
    ).get_cartera("tok")
    assert out["proyecto"] == []
    m = out["mantenimiento"][0]
    assert m["trabajadas"] == 134                # de Time Logs, no de hsTrabajadas
    assert m["diff"] == -6                       # 128 - 134
    assert m["facturadas"] == m["certificadas"] == 128   # Mantenimiento: facturadas = certificadas
    assert m["fecha"] == "2026-08-01"            # primer día del mes
    assert m["cp"] == "María López" and m["vendidasMensuales"] == 140 and m["status"] == "En curso"


def test_service_mantenimiento_rows_carry_anio_mes_for_ux1_facets():
    # UX1: el facet Año/Mes de Mantenimiento necesita estos campos crudos
    # (no derivados de `fecha`) en cada fila para poder filtrar.
    out = ReportingService(_FakeRepo([_mant_fact(anio="2026", mes="Agosto")])).get_cartera("tok")
    m = out["mantenimiento"][0]
    assert m["anio"] == "2026"
    assert m["mes"] == "8"  # numérico, no el nombre largo del mes


def test_service_proyecto_aggregates_over_months_facturadas_not_resummed():
    facts = [
        _proy_fact(mes="Enero", hsTrabajadas=200.0, hsVendidas=220.0, hsCertificadas=190.0),
        _proy_fact(mes="Febrero", hsTrabajadas=140.0, hsVendidas=140.0, hsCertificadas=132.0),
    ]
    out = ReportingService(_FakeRepo(facts)).get_cartera("tok")
    assert out["mantenimiento"] == []
    assert len(out["proyecto"]) == 1
    p = out["proyecto"][0]
    # `certificadas`/`vendidas` SIGUEN agregando por mes de la fuente.
    assert p["certificadas"] == 322 and p["vendidas"] == 360
    assert p["facturadas"] == 290                # XLSX total, NUNCA re-sumado por mes
    # `trabajadas` YA NO sale del campo manual `hsTrabajadas` (200+140=340):
    # su autoridad es Time Logs (OPTION A). Este fixture no aporta Time Logs,
    # así que el resultado correcto es 0.0 — semántica de ausencia, nunca el
    # campo manual de vuelta. La reconciliación real con Time Logs vive en
    # test_reporting_cartera_union_grain.py.
    assert p["trabajadas"] == 0.0
    assert p["diff"] == 322                       # 322 - 0
    assert p["idProyecto"] == "DAN-1042" and p["anioInicio"] == "2026"


def test_service_never_leaks_internal_field_names():
    out = ReportingService(_FakeRepo([_mant_fact(), _proy_fact()])).get_cartera("tok")
    blob = repr(out)
    for internal in ("HsFacturadas", "TipodeProyecto", "EmpresaId", "ProyectoId", "Suma de Cantidad"):
        assert internal not in blob


# --- Router: auth + éxito + fail-soft localizado -----------------------------

from fastapi.testclient import TestClient  # noqa: E402

from webapp.main import app  # noqa: E402
from webapp.routers import reporting as reporting_router  # noqa: E402
from webapp.services import auth_service  # noqa: E402

auth_service.ALLOWED_USERS = "ALL"
auth_service.OPERATIONS_ALLOWED_USERS = "staffing-test@danwor.com,topbar-full@danwor.com,onedanwor-full@danwor.com"
_SESSION = auth_service.create_session(
    {"email": "staffing-test@danwor.com", "name": "Rep Test", "first_name": "Rep", "family_name": "Test"}
)
_client = TestClient(app)
_client.cookies.set(auth_service.COOKIE_NAME, _SESSION)

_OK_CARTERA = {"mantenimiento": [{"cliente": "YPF"}], "proyecto": []}


class _RecordingReportingService:
    def __init__(self, cartera=None, error=None):
        self._cartera = cartera if cartera is not None else _OK_CARTERA
        self._error = error

    def get_cartera(self, access_token):
        if self._error is not None:
            raise self._error
        return self._cartera


def test_cartera_endpoint_requires_operations_session():
    anon = TestClient(app)
    r = anon.post("/operaciones/reporting/integral/data/cartera", data={"graph_token": "tok"}, follow_redirects=False)
    assert r.status_code == 303


def test_cartera_endpoint_success_returns_normalized(monkeypatch):
    monkeypatch.setattr(reporting_router, "ReportingService", lambda: _RecordingReportingService())
    r = _client.post("/operaciones/reporting/integral/data/cartera", data={"graph_token": "tok"})
    assert r.status_code == 200
    body = r.json()
    assert body["state"] == "ok"
    assert body["cartera"]["mantenimiento"][0]["cliente"] == "YPF"


def test_cartera_endpoint_missing_token():
    r = _client.post("/operaciones/reporting/integral/data/cartera", data={"graph_token": ""})
    assert r.json() == {"state": "missing_token"}


def test_cartera_endpoint_success_response_is_only_state_and_cartera(monkeypatch):
    # Cerrada la investigación: la respuesta ok NO expone diagnostics (internal
    # names) — solo state + cartera normalizada.
    monkeypatch.setattr(reporting_router, "ReportingService", lambda: _RecordingReportingService())
    r = _client.post("/operaciones/reporting/integral/data/cartera", data={"graph_token": "tok"})
    body = r.json()
    assert body["state"] == "ok"
    assert set(body.keys()) == {"state", "cartera"}


def test_cartera_endpoint_source_error_is_localized_fail_soft(monkeypatch):
    monkeypatch.setattr(
        reporting_router, "ReportingService",
        lambda: _RecordingReportingService(error=ReportingSourceError(stage="workbook")),
    )
    r = _client.post("/operaciones/reporting/integral/data/cartera", data={"graph_token": "tok"})
    assert r.status_code == 200  # nunca 5xx: no tumba la página
    assert r.json() == {"state": "source_error", "stage": "workbook"}


def test_cartera_endpoint_missing_config_is_config_incomplete(monkeypatch):
    monkeypatch.setattr(
        reporting_router, "ReportingService",
        lambda: _RecordingReportingService(error=ReportingConfigurationError("site no configurado")),
    )
    r = _client.post("/operaciones/reporting/integral/data/cartera", data={"graph_token": "tok"})
    assert r.status_code == 200
    assert r.json() == {"state": "config_incomplete"}


def test_cartera_endpoint_arbitrary_runtime_error_is_not_config_incomplete(monkeypatch):
    # Un RuntimeError genérico (parseo/join/etc.) NUNCA se lee como config
    # incompleta — es un source_error/unexpected, fail-soft, con log server-side.
    monkeypatch.setattr(
        reporting_router, "ReportingService",
        lambda: _RecordingReportingService(error=RuntimeError("boom en el join")),
    )
    r = _client.post("/operaciones/reporting/integral/data/cartera", data={"graph_token": "tok"})
    assert r.status_code == 200
    assert r.json() == {"state": "source_error", "stage": "unexpected"}
    assert "boom en el join" not in r.text  # nunca el texto crudo al browser


def test_cartera_endpoint_graph_error_exposes_safe_stage_status(monkeypatch):
    monkeypatch.setattr(
        reporting_router, "ReportingService",
        lambda: _RecordingReportingService(error=ReportingGraphError(stage="proyectos", status_code=403)),
    )
    r = _client.post("/operaciones/reporting/integral/data/cartera", data={"graph_token": "tok"})
    assert r.status_code == 200
    assert r.json() == {"state": "source_error", "stage": "proyectos", "http_status": 403}


def test_cartera_endpoint_never_echoes_token(monkeypatch):
    secret = "eyJ-super-secret-token"
    monkeypatch.setattr(reporting_router, "ReportingService", lambda: _RecordingReportingService())
    r = _client.post("/operaciones/reporting/integral/data/cartera", data={"graph_token": secret})
    assert secret not in r.text


# --- Router: `perspectiva` (P5, owner live QA de performance) ----------------

class _PerspectivaAwareReportingService:
    """Registra qué método real se invocó — prueba que `perspectiva` en el
    router dispara SOLO el fetch de esa perspectiva, nunca el combo completo
    (que sería el único camino con riesgo de tocar la XLSX de Facturadas
    para Mantenimiento)."""

    def __init__(self):
        self.called = []

    def get_cartera(self, access_token):
        self.called.append("combo")
        return {"mantenimiento": [{"cliente": "combo"}], "proyecto": [{"cliente": "combo"}]}

    def get_cartera_mantenimiento(self, access_token, *, anios=None):
        self.called.append("mantenimiento")
        return [{"cliente": "solo-mantenimiento"}]

    def get_cartera_proyecto(self, access_token, *, cache_scope=None):
        self.called.append("proyecto")
        self.last_cache_scope = cache_scope
        return [{"cliente": "solo-proyecto"}]


def test_cartera_endpoint_perspectiva_mantenimiento_calls_only_that_method(monkeypatch):
    svc = _PerspectivaAwareReportingService()
    monkeypatch.setattr(reporting_router, "ReportingService", lambda: svc)
    r = _client.post(
        "/operaciones/reporting/integral/data/cartera",
        data={"graph_token": "tok", "perspectiva": "mantenimiento"},
    )
    body = r.json()
    assert svc.called == ["mantenimiento"]  # nunca "combo" — nunca toca Proyecto/XLSX
    assert body["cartera"] == {"mantenimiento": [{"cliente": "solo-mantenimiento"}]}
    assert "proyecto" not in body["cartera"]


def test_cartera_endpoint_perspectiva_proyecto_calls_only_that_method(monkeypatch):
    svc = _PerspectivaAwareReportingService()
    monkeypatch.setattr(reporting_router, "ReportingService", lambda: svc)
    r = _client.post(
        "/operaciones/reporting/integral/data/cartera",
        data={"graph_token": "tok", "perspectiva": "proyecto"},
    )
    body = r.json()
    assert svc.called == ["proyecto"]
    assert body["cartera"] == {"proyecto": [{"cliente": "solo-proyecto"}]}


def test_cartera_endpoint_perspectiva_proyecto_passes_session_email_as_cache_scope(monkeypatch):
    # CACHE VERSIONED HISTORICAL TIME LOGS corrective (Fase 2, ver
    # HANDOFF): el router resuelve `cache_scope` desde la sesión Hub YA
    # autenticada (claim `sub` del JWT, ver `_SESSION` arriba) — NUNCA
    # desde `graph_token`. Confirma el wiring exacto.
    svc = _PerspectivaAwareReportingService()
    monkeypatch.setattr(reporting_router, "ReportingService", lambda: svc)
    _client.post(
        "/operaciones/reporting/integral/data/cartera",
        data={"graph_token": "tok", "perspectiva": "proyecto"},
    )
    assert svc.last_cache_scope == "staffing-test@danwor.com"


def test_cartera_endpoint_perspectiva_omitted_preserves_retrocompat_combo(monkeypatch):
    svc = _PerspectivaAwareReportingService()
    monkeypatch.setattr(reporting_router, "ReportingService", lambda: svc)
    r = _client.post("/operaciones/reporting/integral/data/cartera", data={"graph_token": "tok"})
    body = r.json()
    assert svc.called == ["combo"]
    assert body["cartera"]["mantenimiento"] == [{"cliente": "combo"}]
    assert body["cartera"]["proyecto"] == [{"cliente": "combo"}]


# --- Router: `anios` (HOT WINDOW / COLD HISTORY corrective) ------------------

class _AniosRecordingReportingService:
    def __init__(self):
        self.received_anios = "unset"

    def get_cartera_mantenimiento(self, access_token, *, anios=None):
        self.received_anios = anios
        return [{"cliente": "mant"}]

    def get_cartera_mantenimiento_anios(self, access_token):
        return [{"anio": "2024", "estado": "En curso"}, {"anio": "2023", "estado": "Finalizado"}]


def test_cartera_endpoint_anios_param_parsed_as_explicit_list(monkeypatch):
    svc = _AniosRecordingReportingService()
    monkeypatch.setattr(reporting_router, "ReportingService", lambda: svc)
    r = _client.post(
        "/operaciones/reporting/integral/data/cartera",
        data={"graph_token": "tok", "perspectiva": "mantenimiento", "anios": "2024, 2023"},
    )
    assert r.json()["state"] == "ok"
    assert svc.received_anios == ["2024", "2023"]


def test_cartera_endpoint_anios_omitted_lets_service_apply_default(monkeypatch):
    svc = _AniosRecordingReportingService()
    monkeypatch.setattr(reporting_router, "ReportingService", lambda: svc)
    r = _client.post(
        "/operaciones/reporting/integral/data/cartera",
        data={"graph_token": "tok", "perspectiva": "mantenimiento"},
    )
    assert r.json()["state"] == "ok"
    assert svc.received_anios is None  # el router NUNCA calcula el default — lo hace el service


def test_cartera_endpoint_mantenimiento_response_never_bundles_anios_disponibles(monkeypatch):
    # Costo separado a propósito (ver docstring del router) — el camino
    # rápido de Mantenimiento nunca paga la lectura extra.
    svc = _AniosRecordingReportingService()
    monkeypatch.setattr(reporting_router, "ReportingService", lambda: svc)
    r = _client.post(
        "/operaciones/reporting/integral/data/cartera",
        data={"graph_token": "tok", "perspectiva": "mantenimiento"},
    )
    assert "aniosDisponibles" not in r.json()["cartera"]


def test_cartera_anios_disponibles_endpoint_requires_operations_session():
    anon = TestClient(app)
    r = anon.post(
        "/operaciones/reporting/integral/data/cartera/anios-disponibles",
        data={"graph_token": "tok"}, follow_redirects=False,
    )
    assert r.status_code == 303


def test_cartera_anios_disponibles_endpoint_returns_pairs(monkeypatch):
    svc = _AniosRecordingReportingService()
    monkeypatch.setattr(reporting_router, "ReportingService", lambda: svc)
    r = _client.post(
        "/operaciones/reporting/integral/data/cartera/anios-disponibles", data={"graph_token": "tok"},
    )
    body = r.json()
    assert body["state"] == "ok"
    assert body["aniosDisponibles"] == [{"anio": "2024", "estado": "En curso"}, {"anio": "2023", "estado": "Finalizado"}]


def test_cartera_anios_disponibles_endpoint_missing_token():
    r = _client.post("/operaciones/reporting/integral/data/cartera/anios-disponibles", data={"graph_token": ""})
    assert r.json() == {"state": "missing_token"}


# --- Browser wiring (source-seam, sin runtime JS) ----------------------------

from pathlib import Path  # noqa: E402

_REPORTING_JS = (Path(__file__).resolve().parents[2] / "webapp" / "static" / "reporting.js").read_text(encoding="utf-8")
_TOOLS_CSS = (Path(__file__).resolve().parents[2] / "webapp" / "static" / "tools.css").read_text(encoding="utf-8")


def test_cartera_fetches_the_backend_endpoint():
    assert '"/operaciones/reporting/integral/data/cartera"' in _REPORTING_JS
    # El token ya no se lee del storage acá: lo resuelve la autoridad
    # compartida del Hub (`dwGraphToken`, hub_auth.js). Lo que este test
    # sigue afirmando es lo mismo de antes —que Cartera manda un token con
    # su request—, ahora contra el contrato correcto.
    assert "await dwGraphToken()" in _REPORTING_JS
    assert "dw-graph-token" not in _REPORTING_JS  # nunca vuelve a poseerlo


def test_cartera_populates_existing_arrays_in_place_not_reassign():
    # Preserva el motor: muta los arrays que CARTERA_CONFIG referencia.
    # P5 (owner live QA de performance): Mantenimiento/Proyecto tienen
    # fetch/status INDEPENDIENTES — Mantenimiento NUNCA dispara la XLSX.
    mant_block = _REPORTING_JS[_REPORTING_JS.index("async function loadCarteraMantenimiento"):_REPORTING_JS.index("async function loadCarteraProyecto")]
    assert "replaceRows(CARTERA_MANTENIMIENTO_ROWS, data.cartera.mantenimiento)" in mant_block
    assert 'perspectiva: "mantenimiento"' in mant_block

    proy_block = _REPORTING_JS[_REPORTING_JS.index("async function loadCarteraProyecto"):_REPORTING_JS.index("function carteraStatusFor")]
    assert "replaceRows(CARTERA_PROYECTO_ROWS, data.cartera.proyecto)" in proy_block
    assert 'perspectiva: "proyecto"' in proy_block
    assert "state.carteraAnioInicio = carteraDefaultAnioInicio()" in proy_block  # re-deriva default real


def test_cartera_render_is_gated_by_data_status():
    # SHARED REPORTING INITIAL-LOAD CONTRACT (owner corrective, ver HANDOFF):
    # el trigger de hidratación (ambas perspectivas + metadata de años) ya NO
    # vive en `renderCartera()` — se centralizó en el shell
    # (`TAB_HYDRATION_TRIGGERS.cartera`, ver `render()`) para poder mostrar
    # SOLO tabs+loader durante la primera hidratación sin dejar el tab
    # "idle" para siempre. El dispatcher sigue gateado SOLO por el status de
    # la perspectiva ACTIVA (progresivo).
    assert "cartera: hydrateCarteraTab," in _REPORTING_JS
    render_block = _REPORTING_JS[_REPORTING_JS.index("function renderCartera()"):_REPORTING_JS.index("const stateKey")]
    assert "hydrateCarteraTab();" not in render_block
    assert "if (status !== \"ready\") return renderCarteraDataState();" in render_block


def test_hydrate_cartera_tab_loads_only_the_active_perspective_plus_anios_metadata():
    # ACTIVE EXPERIENCE + ACTIVE PERSPECTIVE ONLY corrective (owner decision
    # 2026-08-19 — lifecycle audit): reemplaza el viejo contrato "ambas
    # perspectivas eager" (protegido antes por este mismo test) — Proyecto
    # NUNCA se dispara hasta que sea la perspectiva REALMENTE activa. La
    # metadata de años de Mantenimiento sigue sin condicionar por
    # perspectiva (liviana, y el default de Cartera es siempre
    # Mantenimiento al bootstrap). Cobertura de runtime real (conteo de
    # fetches, idle/ready, revisitas) en
    # test_reporting_cartera_lifecycle_active_perspective_only.py.
    fn = _REPORTING_JS[_REPORTING_JS.index("function hydrateCarteraTab()"):_REPORTING_JS.index("function carteraMantenimientoLoadedYears")]
    assert "const persp = state.perspective.cartera;" in fn
    assert 'if (carteraStatusFor(persp) === "idle") loadCarteraFor(persp);' in fn
    assert 'if (state.carteraMantenimientoAniosStatus === "idle") loadCarteraMantenimientoAnios();' in fn
    # Nunca vuelve a disparar AMBAS perspectivas incondicionalmente.
    assert 'if (state.carteraProyectoStatus === "idle") loadCarteraProyecto();' not in fn
    assert 'if (state.carteraMantenimientoStatus === "idle") loadCarteraMantenimiento();' not in fn


def test_cartera_failure_is_localized_fail_soft():
    mant_block = _REPORTING_JS[_REPORTING_JS.index("async function loadCarteraMantenimiento"):_REPORTING_JS.index("async function loadCarteraProyecto")]
    assert 'state.carteraMantenimientoStatus = "error"' in mant_block
    assert "catch (e)" in mant_block
    # El estado de error de Cartera NO menciona otras experiencias como rotas.
    assert "siguen disponibles" in _REPORTING_JS


def test_filter_bar_hidden_until_first_hydration_only():
    # INITIAL HYDRATION CONTRACT (owner): la filter bar se oculta SOLO durante la
    # PRIMERA hidratación (nunca hidratado). Una vez hidratado el tab, conmutar a
    # una perspectiva hermana en carga NO vuelve a ocultarla (se mantiene + loader
    # contextual en el contenido) — de ahí el guard por reportingTabHydrated,
    # no solo por status de perspectiva.
    fb = _REPORTING_JS[_REPORTING_JS.index("function renderCarteraFilterBar()"):_REPORTING_JS.index("function renderRentabilidadFilterBar()")]
    assert 'if (!state.reportingTabHydrated.cartera && carteraStatusFor(persp) !== "ready") return "";' in fb


# --- UX1: contrato de filtro Año/Mes de Mantenimiento -------------------------

def test_meses_largos_is_fixed_twelve_months_not_derived_from_data():
    idx = _REPORTING_JS.index("const MESES_LARGOS")
    block = _REPORTING_JS[idx:idx + 260]
    for mes in ("Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio",
                "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre"):
        assert mes in block


def test_cartera_mes_field_renders_all_twelve_months_plus_todos():
    # AÑO/MES MULTI-SELECT corrective (DECISIÓN VISUAL OWNER): Mes pasó de
    # <select> a pills — los 12 meses siguen fijos, sin importar data
    # (MESES.map, nunca un array recortado por facet), y selección vacía
    # sigue significando "Todos" (nunca un pill "Todos" explícito, MISMO
    # criterio que Horas Productivas).
    fn = _REPORTING_JS[_REPORTING_JS.index("function carteraMesPills()"):_REPORTING_JS.index("function pillMultiField(")]
    assert "MESES.map((abrev, i) =>" in fn
    assert "MESES_LARGOS[i]" in fn


# --- P2: single-flight cache per source-family (owner live QA de performance) -

def test_all_source_family_loaders_flip_status_synchronously_before_any_await():
    # P2: no hay infraestructura de cache nueva — el single-flight ya lo
    # garantiza el propio patrón idle/loading/ready/error: cada loader pone
    # `status = "loading"` como PRIMERA sentencia, antes de cualquier
    # `await`. Como JS es single-threaded, ese flip ocurre de forma síncrona
    # ANTES de que el loop de eventos pueda ejecutar un segundo render() que
    # vea `status === "idle"` — un segundo disparo del mismo tab nunca
    # dispara un fetch concurrente duplicado.
    loaders = {
        "loadCarteraMantenimiento": "state.carteraMantenimientoStatus = \"loading\";",
        "loadCarteraProyecto": "state.carteraProyectoStatus = \"loading\";",
        "loadRentabilidadData": "state.rentabilidadDataStatus = \"loading\";",
        "loadHorasVendidasData": "state.horasVendidasDataStatus = \"loading\";",
        "loadBatchAData": "state.batchADataStatus = \"loading\";",
        "loadBatchBData": "state.batchBDataStatus = \"loading\";",
    }
    for fn_name, expected_first_line in loaders.items():
        start = _REPORTING_JS.index(f"async function {fn_name}()")
        body_open = _REPORTING_JS.index("{", start) + 1
        first_statement = _REPORTING_JS[body_open:body_open + 200].strip().splitlines()[0].strip()
        assert first_statement == expected_first_line, fn_name
        # Ninguna sentencia `await` aparece ANTES de ese flip.
        pre_flip = _REPORTING_JS[body_open:_REPORTING_JS.index(expected_first_line, body_open)]
        assert "await" not in pre_flip, fn_name


def test_facet_config_cartera_anio_and_mes_apply_on_base_rows_not_predicate():
    # AÑO/MES MULTI-SELECT corrective (DECISIÓN VISUAL OWNER): Año/Mes
    # dejaron el motor de `fields`/`predicate` (valor único) — son Sets
    # multi-select aplicados directo sobre `baseRows`, MISMO criterio que
    # Año de Inicio/Estado. Año conserva su PROPIA reconciliación
    # metadata-driven (`pruneCarteraAnio`, ver
    # test_stale_anio_is_pruned_via_own_metadata_driven_function abajo).
    idx = _REPORTING_JS.index("const FACET_CONFIG = {")
    idx = _REPORTING_JS.index("cartera:", idx)
    block = _REPORTING_JS[idx:idx + 2500]
    fields_idx = block.index("fields:")
    base_rows_block = block[:fields_idx]
    fields_block = block[fields_idx:fields_idx + 200]
    assert "state.carteraAnios" in base_rows_block
    assert "state.carteraMeses" in base_rows_block
    assert 'key: "anio"' not in fields_block
    assert 'key: "mes"' not in fields_block


def test_stale_anio_is_pruned_via_own_metadata_driven_function():
    assert "function carteraAnioOptions()" in _REPORTING_JS
    assert "function pruneCarteraAnio()" in _REPORTING_JS
    fn = _REPORTING_JS[_REPORTING_JS.index("function carteraAnioOptions()"):_REPORTING_JS.index("function pruneCarteraAnio()")]
    # Unión metadata (años fríos incluidos) + filas ya cargadas, faceteado por
    # Estado — vía `_statusOrAllIfEmpty` (ESTADO VACÍO corrective: mismo
    # gate, pero nunca colapsa las opciones a nada si Estado queda sin
    # ningún tildado).
    assert "carteraMantenimientoAniosDisponibles" in fn
    assert "CARTERA_MANTENIMIENTO_ROWS" in fn
    assert "_statusOrAllIfEmpty(state.carteraStatus)" in fn
    prune = _REPORTING_JS[_REPORTING_JS.index("function pruneCarteraAnio()"):]
    prune = prune[:prune.index("\n}")]
    # Año pasó a multi-select (Set) — poda individual de años obsoletos,
    # nunca vacía toda la selección por uno solo inválido.
    assert "available.has(y)" in prune
    assert "state.carteraAnios = new Set(" in prune


def test_estado_vacio_keeps_temporal_pills_structure_disabled_not_collapsed():
    # ESTADO VACÍO corrective (owner — mejora transversal): con AMBOS
    # checkboxes de Estado destildados, Año/Mes NUNCA colapsan a "sin
    # opciones" — conservan la MISMA estructura de pills (calculada como si
    # Estado tuviera todo tildado) pero el grupo entero queda deshabilitado
    # (`disabledAll`) hasta que el usuario tilde al menos un Estado.
    assert "function _statusHasAnySelected(statusMap)" in _REPORTING_JS
    assert "function _statusOrAllIfEmpty(statusMap)" in _REPORTING_JS
    gate = _REPORTING_JS[_REPORTING_JS.index("function _statusOrAllIfEmpty"):_REPORTING_JS.index("function _statusOrAllIfEmpty") + 400]
    assert "forced[k] = true" in gate
    pill_field = _REPORTING_JS[_REPORTING_JS.index("function pillMultiField"):_REPORTING_JS.index("function checkboxGroupField")]
    assert "opts.disabledAll" in pill_field
    for fn_name, status_var in (
        ("carteraAnioPills", "carteraStatus"), ("carteraMesPills", "carteraStatus"),
        ("rentabilidadAnioPills", "rentabilidadStatus"), ("rentabilidadMesPills", "rentabilidadStatus"),
        ("horasVendidasAnioPills", "horasVendidasStatus"), ("horasVendidasMesPills", "horasVendidasStatus"),
    ):
        start = _REPORTING_JS.index(f"function {fn_name}(")
        body = _REPORTING_JS[start:_REPORTING_JS.index("\n}", start)]
        assert f"_statusHasAnySelected(state.{status_var})" in body, fn_name
        assert "disabledAll" in body, fn_name
    css = (Path(__file__).resolve().parents[2] / "webapp" / "static" / "tools.css").read_text(encoding="utf-8")
    assert ".dashboard-view-switch button:disabled { opacity: .35; cursor: not-allowed; }" in css


# ============================================================================
# Phase 0 — hardening transversal (loader / geometría de matriz / formateo KPI)
# ============================================================================

_REPORTING_CSS = (Path(__file__).resolve().parents[2] / "webapp" / "static" / "reporting.css").read_text(encoding="utf-8")


def test_loader_reuses_shared_hub_primitive_not_a_reporting_copy():
    # 0A: Reporting reusa el MISMO primitivo de Asignaciones, no un empty-panel.
    assert 'reportingLoadingHtml("Cargando datos de Cartera' in _REPORTING_JS
    helper = _REPORTING_JS[_REPORTING_JS.index("function reportingLoadingHtml"):_REPORTING_JS.index("function reportingSourceErrorHtml")]
    assert "section-analysis-loading" in helper and "section-analysis-spinner" in helper


def test_semantic_number_formatter_exists_and_kpis_use_it():
    # 0C: fmtNum semántico. CRITICAL HUB STANDARDIZATION CORRECTIVE (owner
    # live QA — raw float leak): los tiles KPI YA NO construyen su propio
    # markup con `fmtNum(...)` inline — pasan por el primitivo único
    # `renderDashboardKpi`, ver test_reporting_kpi_primitive.py para el
    # contrato completo de esa costura.
    assert "function fmtNum(value, type" in _REPORTING_JS
    assert "Intl.NumberFormat" in _REPORTING_JS
    assert "function renderDashboardKpi(" in _REPORTING_JS
    # diffCell usa el formateo con signo (sin ruido de punto flotante).
    diff = _REPORTING_JS[_REPORTING_JS.index("function diffCell"):_REPORTING_JS.index("function diffPlain")]
    assert 'fmtNum(diff, "signed")' in diff


def test_matrix_geometry_is_fixed_layout_for_stable_columns():
    # 0B, ahora del primitivo compartido `.dashboard-table` (Phase S, ver
    # HANDOFF): table-layout: fixed ⇒ expand/collapse no mueve columnas.
    # Reporting consume la clase compartida (reporting.js) + su propio
    # override de ancho (reporting.css, min-width de la matriz mensual).
    assert "table.dashboard-table { width: 100%; border-collapse: collapse; font-size: .76rem; table-layout: fixed; }" in _TOOLS_CSS
    assert 'class="dashboard-table"' in _REPORTING_JS
    assert ".reporting-integral-preview table.dashboard-table { min-width: 900px; }" in _REPORTING_CSS


def test_kpi_tile_has_overflow_safe_styles():
    # 0C, ahora del primitivo compartido `.activity-kpi`/`.activity-kpi__value`
    # (Phase S, ver HANDOFF): tiles no se solapan con valores de 5–7 dígitos
    # en NINGÚN consumidor (Asignaciones/Reporting/futuro Activity), no solo
    # Reporting.
    assert ".activity-kpi { min-width: 0;" in _TOOLS_CSS
    assert "white-space: nowrap; font-variant-numeric: tabular-nums" in _TOOLS_CSS


# ============================================================================
# Q1 — Rentabilidad / Certificadas vs Ejecutadas / Detalle por proyecto
# (owner live QA: Tipo/Proyecto comprimidas, columnas numéricas colisionan)
# ============================================================================

def test_rentabilidad_cert_ejec_detail_uses_compact_role_not_a_pixel_hack():
    # Tipo (categórico corto) usa el rol GENÉRICO `--compact` del vocabulario
    # compartido — Proyecto (descriptivo) sigue `.left` sin ancho fijo, así
    # absorbe el espacio remanente en vez de repartirlo en partes iguales con
    # Tipo. Nunca un ancho puntual/inline de Rentabilidad.
    detail_start = _REPORTING_JS.index("function renderRentabilidadCertEjecutadas")
    detail_body = _REPORTING_JS[detail_start:_REPORTING_JS.index("function renderRentabilidad()")]
    assert '<th class="left dashboard-table-col--compact">Tipo</th><th class="left">Proyecto</th>' in detail_body
    assert "style=" not in detail_body[detail_body.index("<thead>"):detail_body.index("</thead>") + 8]
    assert "table.dashboard-table th.dashboard-table-col--compact { width: 100px; }" in _TOOLS_CSS


def test_dashboard_table_header_wraps_instead_of_colliding():
    # Un header numérico más largo que su columna fija de 88px (ej. "Margen
    # Relativo") ya no desborda sobre la columna vecina — el primitivo
    # compartido permite wrap a 2 líneas (geometría estable, sin JS de
    # medición). Cambio TRANSVERSAL en `.dashboard-table`, no un ajuste
    # puntual de Rentabilidad.
    assert "white-space: normal; line-height: 1.15;" in _TOOLS_CSS
    header_rule = _TOOLS_CSS[_TOOLS_CSS.index("table.dashboard-table th {"):_TOOLS_CSS.index("table.dashboard-table th {") + 400]
    assert "white-space: nowrap;" not in header_rule


# ============================================================================
# Corrective pre-Batch-A (owner live QA — Cartera es la plantilla canónica)
# ============================================================================

def test_cartera_estado_terminal_column_uses_shared_status_modifier():
    # S.1: Estado es la ÚLTIMA columna en ambas perspectivas de Cartera — usa
    # el primitivo compartido `.dashboard-table-col--status`/`--cell--status`
    # (tools.css), NUNCA un parche visual de un solo uso
    # (`.reporting-status-last` o similar). El modificador es OPCIONAL
    # (`c.statusTerminal`) — Asignaciones (Estado en el medio de su tabla) no
    # lo usa y sigue con el comportamiento `.left` default.
    assert "statusTerminal: true" in _REPORTING_JS
    assert '"dashboard-table-col--status" : (c.numeric ? "" : "left")' in _REPORTING_JS
    assert '"dashboard-table-cell--status" : (c.numeric ? "" : "left")' in _REPORTING_JS
    assert "table.dashboard-table th.dashboard-table-col--status { text-align: right;" in _TOOLS_CSS
    assert "table.dashboard-table td.dashboard-table-cell--status { text-align: right; }" in _TOOLS_CSS
    assert ".reporting-status-last" not in _TOOLS_CSS and ".reporting-status-last" not in _REPORTING_CSS
    # Ancho estándar de columna numérica restaurado (regresión de Phase S:
    # se perdió al fusionar tree-matrix/data-table) — sin esto, Estado
    # terminal comparte el ancho ancho de las columnas de medida.
    assert "table.dashboard-table th:not(.left):not(.row-label-col) { width: 88px; }" in _TOOLS_CSS


def test_cartera_asignaciones_status_column_placement_is_not_globally_forced():
    # S.1 (regla explícita del owner): el modificador de status compacto/
    # derecha NUNCA se fuerza globalmente — Asignaciones sigue usando
    # `class="left"` para su columna Estado (en el medio de la tabla, no
    # terminal), no `dashboard-table-col--status`.
    from pathlib import Path
    staffing_html = (
        Path(__file__).resolve().parents[2] / "webapp" / "templates" / "operations" / "staffing_dashboard.html"
    ).read_text(encoding="utf-8")
    assert '<th class="left">Estado</th>' in staffing_html
    assert "dashboard-table-col--status" not in staffing_html


def test_cartera_proyecto_hierarchy_shows_business_description_not_id():
    # S.2: la jerarquía de Cartera/Proyecto YA NO tiene a `idProyecto` como
    # nivel visible (mostraba el ID técnico, ej. "PROJ-0206", en vez de la
    # descripción de negocio que el filtro Proyecto sí muestra) — Cliente ›
    # Proyecto (2 niveles). `idProyecto` sigue en cada fila como campo
    # interno (join key / desambiguación), nunca se descartó del dato.
    proyecto_block = _REPORTING_JS[_REPORTING_JS.index('"Proyecto": {'):_REPORTING_JS.index('"Proyecto": {') + 900]
    assert 'keys: ["cliente", "proyecto"]' in proyecto_block
    assert 'labels: ["Cliente", "Proyecto"]' in proyecto_block
    assert '"idProyecto"' not in proyecto_block.split("keys:")[1].split("\n")[0]
    # El facet/filtro de Proyecto sigue usando el mismo campo `proyecto`
    # (descripción de negocio) que la jerarquía — dropdown y tabla nunca
    # divergen.
    assert 'facetSelectField("cartera", "proyecto", "Proyecto")' in _REPORTING_JS


def test_cartera_diff_hs_kpi_recomputes_from_filtered_leaf_rows_not_double_formatted():
    # S.3: `carteraKpiTiles` NO debe pasar `col.rollup(rows)` (HTML de
    # `diffCell`, pensado para la celda de tabla) por un formateador de
    # nuevo — eso da NaN → "—" aunque el detalle filtrado tenga valores
    # reales. El tile recalcula el número CRUDO con el mismo `sumField`
    # aditivo (SUM(diffs) == SUM(Certificadas) − SUM(Trabajadas)) y se lo
    # pasa, SIN formatear, al primitivo `renderDashboardKpi` (CRITICAL HUB
    # STANDARDIZATION CORRECTIVE — el formateo es responsabilidad EXCLUSIVA
    # del primitivo, nunca del call site, ver
    # test_reporting_kpi_primitive.py).
    # Delimitado por el CIERRE real de la función, no por un largo mágico: un
    # comentario nuevo adentro corría el call site fuera de la ventana fija.
    tiles_fn = _REPORTING_JS.split("function carteraKpiTiles", 1)[1].split("\n}", 1)[0]
    assert 'const value = k.key === "diff" ? sumField("diff")(rows) : col.rollup(rows);' in tiles_fn
    assert 'renderDashboardKpi({ label: k.label, value, format: k.key === "diff" ? "signed" : "hours", info })' in tiles_fn
    # Patrón viejo (formateo inline en el call site) no debe quedar.
    assert "fmtNum(sumField(" not in tiles_fn
    assert "fmtNum(col.rollup(rows)" not in tiles_fn


def test_hierarchy_disclosure_is_one_shared_helper_not_duplicated_inline():
    # S.4: UN solo helper (`hierarchyToggleHtml`) genera el control de
    # disclosure para AMBOS renderers de árbol (renderTreeRows/
    # renderBusinessTreeRows) — nunca dos bloques de markup casi-idénticos
    # copiados a mano (el estado previo a este corrective).
    assert _REPORTING_JS.count("function hierarchyToggleHtml") == 1
    assert _REPORTING_JS.count("hierarchyToggleHtml(stateKey, node.pathKey, node.label, hasChildren, isOpen, caretChar)") == 2
    # El markup inline viejo (dos <span class="dashboard-table-toggle"> con
    # el mismo onclick repetido) no debe quedar.
    assert '<span class="dashboard-table-toggle" ${hasChildren' not in _REPORTING_JS


def test_hierarchy_disclosure_control_has_children_is_a_real_accessible_button():
    # S.4 (owner live QA): clic en el triángulo O en el label deben disparar
    # el MISMO toggle — un solo control, no dos manejadores separados que
    # puedan doble-disparar. `<button>` real ⇒ foco/teclado (Enter/Space)
    # gratis del browser, `aria-expanded` sincronizado con el estado.
    helper_fn = _REPORTING_JS[_REPORTING_JS.index("function hierarchyToggleHtml"):_REPORTING_JS.index("function renderTreeRows")]
    assert '<button type="button" class="dashboard-table-toggle"' in helper_fn
    assert "onclick=\"toggleTreeNode(" in helper_fn
    assert 'aria-expanded="${isOpen}"' in helper_fn


def test_hierarchy_leaf_rows_are_not_falsely_expandable():
    # S.4 (owner live QA: "leaf rows must not appear expandable"): sin
    # hijos, el control es un `<span>` NO interactivo — nunca un `<button>`,
    # nunca un `onclick`, nunca affordance de puntero engañosa.
    helper_fn = _REPORTING_JS[_REPORTING_JS.index("function hierarchyToggleHtml"):_REPORTING_JS.index("function renderTreeRows")]
    # Hoja = <span> (no <button>); tolera el atributo `title` agregado después
    # de la clase (owner live QA — tooltip de nombre completo).
    assert '<span class="dashboard-table-toggle dashboard-table-toggle--leaf"' in helper_fn
    assert ".dashboard-table-toggle.dashboard-table-toggle--leaf { cursor: default; }" in _TOOLS_CSS
    assert "button.dashboard-table-toggle { cursor: pointer; }" in _TOOLS_CSS


def test_fmt_num_distinguishes_zero_from_missing():
    # Validación explícita del owner: 0 es un resultado numérico válido
    # (debe renderizar "0"), el em dash es SOLO para null/undefined/"".
    fmt_num_fn = _REPORTING_JS[_REPORTING_JS.index("function fmtNum"):_REPORTING_JS.index("function fmtNum") + 400]
    assert 'if (value === null || value === undefined || value === "") return "—";' in fmt_num_fn
    # 0 no cae en ninguna de esas tres ramas — sigue de largo hasta el
    # formateo real (Number(0) es finito, nunca dispara el fallback "—").


# ============================================================================
# Año de Inicio facet respeta Estado + otros filtros (corrective multidireccional)
# ============================================================================

def test_anio_inicio_options_are_faceted_through_the_shared_engine():
    block = _REPORTING_JS[_REPORTING_JS.index("function carteraAnioInicioOptions"):_REPORTING_JS.index("function carteraDefaultAnioInicio")]
    # Reusa el motor faceteado (facetRowsExcept) sobre las filas ya filtradas por
    # Estado (carteraStatus) — NUNCA distinctValues sobre las filas crudas.
    assert "facetRowsExcept(statusRows, state.filters.cartera" in block
    assert "CARTERA_PROYECTO_ROWS.filter(r => state.carteraStatus[r.status])" in block
    assert "distinctValues(CARTERA_PROYECTO_ROWS" not in block  # ya no ignora Estado


def test_stale_anio_inicio_is_pruned_on_facet_change():
    # Al cambiar Estado u otro facet, se reconcilia la selección obsoleta de año.
    assert "function pruneCarteraAnioInicio()" in _REPORTING_JS
    reconcile = _REPORTING_JS[_REPORTING_JS.index("function reconcileExperienceFacets"):]
    reconcile = reconcile[:reconcile.index("\n}")]
    assert "pruneCarteraAnioInicio();" in reconcile
    assert "pruneCarteraAnio();" in reconcile  # HOT WINDOW/COLD HISTORY: misma reconciliación, Año de Mantenimiento
    prune = _REPORTING_JS[_REPORTING_JS.index("function pruneCarteraAnioInicio"):]
    prune = prune[:prune.index("\n}")]
    assert "available.has(y)" in prune  # conserva solo los válidos, nunca reselecciona


# ============================================================================
# Rentabilidad Año/Mes/Estado migración — DOM + faceting (end-to-end vertical)
# ============================================================================

def _rent_filter_bar():
    return _REPORTING_JS[_REPORTING_JS.index("function renderRentabilidadFilterBar"):_REPORTING_JS.index("function renderHorasVendidasFilterBar")]


def test_rentabilidad_dom_shows_anio_mes_estado_not_desde_hasta():
    # AÑO/MES MULTI-SELECT corrective (propagación del patrón Cartera):
    # Año/Mes son pills, no <select>.
    fb = _rent_filter_bar()
    assert "rentabilidadAnioPills()" in fb and "rentabilidadMesPills()" in fb
    assert "toggleRentabilidadStatus('En curso')" in fb and "toggleRentabilidadStatus('Finalizado')" in fb
    assert 'dateField("rentabilidad"' not in fb          # Desde/Hasta removidos del DOM
    assert '"desde"' not in fb and '"hasta"' not in fb


def test_rentabilidad_estado_default_en_curso_only():
    st = _REPORTING_JS[_REPORTING_JS.index("rentabilidadStatus: {"):_REPORTING_JS.index("rentabilidadStatus: {") + 80]
    assert '"En curso": true' in st and '"Finalizado": false' in st


def test_rentabilidad_anio_options_faceted_by_estado_and_selects():
    block = _REPORTING_JS[_REPORTING_JS.index("function rentabilidadAnioOptions"):_REPORTING_JS.index("function rentabilidadDefaultAnio")]
    # ESTADO VACÍO corrective: gateado vía `_statusOrAllIfEmpty`, no
    # directo contra `state.rentabilidadStatus`.
    assert "_statusOrAllIfEmpty(state.rentabilidadStatus)" in block
    assert "facetRowsExcept(statusRows, state.filters.rentabilidad" in block


def test_rentabilidad_default_anio_current_or_latest():
    # YEAR ORDERING corrective (owner decision 2026-08-19): el default ya
    # NO indexa por posición del array (`options[options.length - 1]`,
    # correcto SOLO mientras las opciones fueran ascendentes) — usa
    # `mostRecentValidYear`, que compara valor numérico y es indiferente
    # al orden de presentación (ver test_most_recent_valid_year_* abajo).
    block = _REPORTING_JS[_REPORTING_JS.index("function rentabilidadDefaultAnio"):_REPORTING_JS.index("function pruneRentabilidadAnio")]
    assert "new Date().getFullYear()" in block and "mostRecentValidYear(options)" in block
    assert "options[options.length - 1]" not in block


def test_rentabilidad_one_canonical_population():
    block = _REPORTING_JS[_REPORTING_JS.index("function rentabilidadFilteredProyectos"):_REPORTING_JS.index("function rentabilidadFilteredProyectos") + 240]
    assert "return rentabilidadSelectedRows();" in block


def test_rentabilidad_cert_ejec_two_panel_shared_grid_and_formatter():
    block = _REPORTING_JS[_REPORTING_JS.index("function renderRentabilidadCertEjecutadas"):_REPORTING_JS.index("function renderRentabilidad()")]
    assert 'class="activity-grid-2"' in block                       # 2-col compartido (Mantenimiento|Proyecto)
    assert 'fmtNum(g.certificadas, "hours")' in block and 'fmtNum(g.trabajadas, "hours")' in block
    assert 'fmtNum(g.margenRelativo, "percent")' in block


def test_anio_mes_fields_are_narrow_everywhere():
    # AÑO/MES MULTI-SELECT corrective (propagación del patrón Cartera a
    # Rentabilidad/Horas Vendidas, ver HANDOFF): las 3 experiencias migraron
    # de <select> a pills — Año se dimensiona a sus propias opciones (nunca
    # `--narrow`, eso era del <select> retirado), Mes usa `--fill` en las 3.
    assert ".reporting-integral-preview .whatsnew-filter--narrow select { min-width: 0; }" in _REPORTING_CSS
    for fn_name in ("carteraMesPills", "rentabilidadMesPills", "horasVendidasMesPills"):
        block = _REPORTING_JS[_REPORTING_JS.index(f"function {fn_name}("):_REPORTING_JS.index(f"function {fn_name}(") + 600]
        assert "fill: true" in block, fn_name


# ============================================================================
# Horas Vendidas Año/Mes/Estado migración — DOM + faceting
# ============================================================================

def test_horas_vendidas_dom_shows_anio_mes_estado_not_desde_hasta():
    fb = _REPORTING_JS[_REPORTING_JS.index("function renderHorasVendidasFilterBar"):_REPORTING_JS.index("// Client Partner primero")]
    assert "horasVendidasAnioPills()" in fb and "horasVendidasMesPills()" in fb
    assert "toggleHorasVendidasStatus('En curso')" in fb and "toggleHorasVendidasStatus('Finalizado')" in fb
    assert 'dateField("horasvendidas"' not in fb
    assert '"desde"' not in fb and '"hasta"' not in fb


def test_horas_vendidas_estado_default_en_curso_only():
    st = _REPORTING_JS[_REPORTING_JS.index("horasVendidasStatus: {"):_REPORTING_JS.index("horasVendidasStatus: {") + 80]
    assert '"En curso": true' in st and '"Finalizado": false' in st


def test_horas_vendidas_anio_options_faceted_and_selected_population():
    opt = _REPORTING_JS[_REPORTING_JS.index("function horasVendidasAnioOptions"):_REPORTING_JS.index("function horasVendidasDefaultAnio")]
    # ESTADO VACÍO corrective: gateado vía `_statusOrAllIfEmpty`.
    assert "_statusOrAllIfEmpty(state.horasVendidasStatus)" in opt
    assert "facetRowsExcept(statusRows, state.filters.horasvendidas" in opt
    sel = _REPORTING_JS[_REPORTING_JS.index("function horasVendidasSelectedRows"):_REPORTING_JS.index("function horasVendidasSelectedRows") + 700]
    # AÑO/MES MULTI-SELECT corrective: `_yearRangeMonthly` (multi-año, OR
    # dentro de dimensión) + `_maskMeses` (Set), MISMO criterio que
    # Rentabilidad.
    assert "vendidasByYear" in sel and "_yearRangeMonthly(" in sel and "_maskMeses(arr, state.horasVendidasMeses)" in sel


def test_matrix_row_label_col_reserves_width():
    # Owner live QA (Horas Vendidas: "no identifica textos para algunos
    # proyectos"). Causa raíz: con `table-layout:fixed` la columna de label
    # jerárquico tomaba solo el ancho RESIDUAL tras las 12 columnas de mes y
    # colapsaba hacia 0px, recortando el nombre de proyecto entero a "…".
    # El fix reserva un ancho fijo a `th.row-label-col` (la matriz scrollea
    # horizontalmente en vez de matar el label).
    assert "table.dashboard-table th.row-label-col { width:" in _TOOLS_CSS


def test_hierarchy_toggle_has_title_tooltip_for_full_label():
    # Red para la cola de nombres larguísimos que aún exceden el ancho
    # reservado: `hierarchyToggleHtml` emite `title=<nombre completo>` (escapado
    # también para la comilla doble, porque `esc` sólo cubre &<>).
    fn = _REPORTING_JS[_REPORTING_JS.index("function hierarchyToggleHtml"):_REPORTING_JS.index("function renderTreeRows")]
    assert 'title="${esc(label).replace(/"/g, "&quot;")}"' in fn
    # El tooltip va TANTO en la hoja (span) como en el nodo con hijos (button).
    assert fn.count("titleAttr") >= 2


# ============================================================================
# Horas Vendidas — scope temporal desacoplado (owner directive: batch-a como
# frontera de runtime; visitar HV no carga Participación)
# ============================================================================

def test_horas_vendidas_has_its_own_scoped_lifecycle_not_batch_a():
    fn = _REPORTING_JS[_REPORTING_JS.index("function renderHorasVendidas() {"):_REPORTING_JS.index("function renderComercialParticipacion")]
    # HV se hidrata por su ciclo PROPIO (ventana temporal), nunca por batch-a.
    # SHARED REPORTING INITIAL-LOAD CONTRACT: el trigger se centralizó en el
    # shell (`TAB_HYDRATION_TRIGGERS.horasvendidas`, ver `render()`) — el
    # dispatcher solo gatea por status.
    assert "horasvendidas: hydrateHorasVendidasTab," in _REPORTING_JS
    assert "hydrateHorasVendidasTab()" not in fn
    assert 'state.horasVendidasDataStatus' in fn
    assert "loadBatchAData()" not in fn
    assert "state.batchADataStatus" not in fn


def test_horas_vendidas_hot_window_loader_hits_scoped_endpoint():
    fn = _REPORTING_JS[_REPORTING_JS.index("async function loadHorasVendidasData"):_REPORTING_JS.index("async function loadHorasVendidasAnios")]
    assert '"/operaciones/reporting/integral/data/horas-vendidas"' in fn
    assert "replaceRows(HORASVENDIDAS_ROWS, data.horasVendidas)" in fn
    # anios NO viaja en la hidratación de la ventana caliente (default server-side).
    assert "anios:" not in fn


def test_horas_vendidas_cold_year_lazy_loads_and_merges_by_project():
    fn = _REPORTING_JS[_REPORTING_JS.index("async function loadHorasVendidasColdYear"):_REPORTING_JS.index("function mergeHorasVendidasRows")]
    # Single-flight por año + POST con el año explícito + merge (no replace).
    assert "state.horasVendidasYearLoading.has(anio)" in fn
    assert "graph_token: token, anios: anio" in fn
    assert "mergeHorasVendidasRows(data.horasVendidas)" in fn
    # El merge fusiona vendidasByYear por proyecto (nunca duplica la fila).
    merge = _REPORTING_JS[_REPORTING_JS.index("function mergeHorasVendidasRows"):_REPORTING_JS.index("function renderHorasVendidasDataState")]
    assert "vendidasByYear" in merge and "HORASVENDIDAS_ROWS.push(inc)" in merge


def test_batch_a_loader_no_longer_populates_horas_vendidas_or_participacion():
    # Frontera de runtime: batch-a (Calendario/Salud) ya NO puebla
    # HORASVENDIDAS_ROWS ni COMERCIAL_PARTICIPACION_ROWS — HV/Participación
    # lo hacen por su propio loader scoped (loadHorasVendidasData/
    # loadComercialParticipacionData).
    fn = _REPORTING_JS[_REPORTING_JS.index("async function loadBatchAData"):_REPORTING_JS.index("function renderBatchADataState")]
    assert "replaceRows(HORASVENDIDAS_ROWS" not in fn
    assert "replaceRows(COMERCIAL_PARTICIPACION_ROWS" not in fn
    assert "replaceRows(CALENDARIO_ROWS, data.calendario)" in fn
    assert "replaceRows(COMERCIAL_SALUD_ROWS, data.salud)" in fn


def test_horas_vendidas_anio_options_offer_cold_metadata_years():
    fn = _REPORTING_JS[_REPORTING_JS.index("function horasVendidasAnioOptions"):_REPORTING_JS.index("function horasVendidasDefaultAnio")]
    # Los años fríos (metadata) se ofrecen en los pills, respetando Estado
    # (vía `_statusOrAllIfEmpty` — ESTADO VACÍO corrective).
    assert "state.horasVendidasAniosDisponibles" in fn
    assert "statusGate[p.estado]" in fn
